"""Synthetic school teams exercise selection, attendance, ranking and workbook output."""
from copy import deepcopy

import openpyxl
import pytest

from exporters.season_results import build_qualified_schools_xlsx
from parsers.season_results import parse_season_results
from parsers.season_results.models import CHAMPIONSHIP, INTERPROVINCIAL, LEAGUE
from services.school_reports import qualified_schools_report, school_type_from_name
from tests.test_season_results import excel_fixture, pdf_fixture


def school_snapshot(school="Example Ps", ages=(9, 11, 13, 15, 11, 13)):
    events = [{"id": i, "name": f"League {i}" if i < 4 else "GN Interprovincial", "event_date": f"2026-08-0{i}",
               "competition_type": LEAGUE if i < 4 else INTERPROVINCIAL} for i in range(1, 5)]
    athletes, results = [], []
    for i, age in enumerate(ages, 1):
        category = f"U/{age} {'GIRLS' if i % 2 else 'BOYS'}" if isinstance(age, int) else age
        athletes.append({"id": i, "athlete_number": f"00{i}", "athlete_name": f"Athlete {i}", "school": school,
                         "category": category, "affiliated": i % 2})
        for event in events:
            results.append({"athlete_id": i, "athlete_number": f"00{i}", "athlete_name": f"Athlete {i}",
                            "school": school if event["id"] < 4 else "", "category": category, "event_id": event["id"],
                            "event_name": event["name"], "event_date": event["event_date"], "competition_type": event["competition_type"],
                            "total_points": str(2500 - i * 100 + (50 if event["id"] == (i % 4) + 1 else 0)),
                            "run_time": "01:20.00", "swim_time": "00:40.00", "status": ""})
    return {"events": events, "athletes": athletes, "results": results}


@pytest.mark.parametrize("name,expected", [("Alpha Ps", "Primary"), ("Alpha HS ", "High"), ("Alpha ps", "Primary"), ("Alpha", "Unassigned"), ("Perhaps", "Unassigned")])
def test_school_type(name, expected):
    assert school_type_from_name(name) == expected


@pytest.mark.parametrize("school,ages", [("Example Ps", (9, 11, 13, 15, 11, 13)), ("Example Hs", (15, 17, 19, 15, 17, 19))])
def test_complete_teams_and_cross_event_best_scores(school, ages):
    snapshot = school_snapshot(school, ages)
    before = deepcopy(snapshot)
    summary, headers, details = qualified_schools_report(snapshot)
    assert summary[0]["Qualified"] == "Yes"
    assert summary[0]["Total school points"] == sum(2550 - i * 100 for i in range(1, 7))
    assert len({r["Athlete number"] for r in details if r["Selected"] == "Yes"}) == 6
    assert len({r["Best event"] for r in details}) == 4
    assert "Affiliated" in headers
    assert len([h for h in headers if h.endswith("]")]) == 4
    assert all(len([h for h in row if h.endswith("]")]) == 4 for row in details)
    assert snapshot == before


def test_missing_mandatory_still_counts_primary_extras():
    summary, _, details = qualified_schools_report(school_snapshot(ages=(11, 13, 15, 11, 13, 15)))
    assert summary[0]["Qualified"] == "No"
    assert summary[0]["Extra places filled"] == "2/2"
    assert "Under-9" in summary[0]["Missing to qualify"]
    assert summary[0]["Total school points"] == 2450 + 2350 + 2250 + 2150 + 2050
    assert sum(r["Selected"] == "Yes" for r in details) == 5


def test_qualified_replacement_preferred_over_higher_score():
    snapshot = school_snapshot(ages=(9, 11, 13, 15, 11, 13, 9))
    for result in snapshot["results"]:
        if result["athlete_id"] == 7:
            result["total_points"] = "1900"
    snapshot["results"] = [r for r in snapshot["results"] if not (r["athlete_id"] == 1 and r["event_id"] == 4)]
    summary, _, rows = qualified_schools_report(snapshot)
    selected = {r["Athlete number"] for r in rows if r["Selected"] == "Yes"}
    assert "001" not in selected and "007" in selected
    assert summary[0]["Qualified"] == "Yes"


def test_unqualified_fallback_scores_and_shortfalls():
    snapshot = school_snapshot()
    snapshot["results"][0]["status"] = "DNF"
    snapshot["results"][1]["run_time"] = "00:00.00"
    summary, _, rows = qualified_schools_report(snapshot)
    assert summary[0]["Qualified"] == "No"
    athlete = next(r for r in rows if r["Athlete number"] == "001")
    assert athlete["Completed league/championship meets"] == 1
    assert "2 more completed" in athlete["Missing to qualify"]
    assert summary[0]["Total school points"] > 0


def test_special_needs_cap_and_maximum_one():
    snapshot = school_snapshot(ages=(9, 11, 13, 15, "SPECIAL NEEDS FEMALE", "SPECIAL NEEDS MALE", 13))
    for r in snapshot["results"]:
        if r["athlete_id"] in {5, 6}:
            r["total_points"] = "3500"
        if r["athlete_id"] == 7:
            r["total_points"] = "1900"
    summary, _, rows = qualified_schools_report(snapshot)
    selected_special = [r for r in rows if r["Selected"] == "Yes" and "SPECIAL" in r["Age group"]]
    assert len(selected_special) == 1
    assert selected_special[0]["School contribution points"] == 2000
    assert selected_special[0]["Best total points"] == 3500
    assert summary[0]["Qualified"] == "Yes"
    assert sum(r["School contribution points"] for r in rows) == summary[0]["Total school points"]


def test_special_needs_cannot_replace_mandatory_or_second_extra():
    snapshot = school_snapshot(ages=(9, 11, 13, 15, "SPECIAL NEEDS FEMALE", "SPECIAL NEEDS MALE"))
    summary, _, _ = qualified_schools_report(snapshot)
    assert summary[0]["Extra places filled"] == "1/2"
    assert summary[0]["Qualified"] == "No"
    snapshot["results"] = [r for r in snapshot["results"] if r["athlete_id"] != 1]
    summary, _, _ = qualified_schools_report(snapshot)
    assert summary[0]["Extra places filled"] == "1/2"


def test_championship_detected_by_name_requires_four_distinct_events():
    snapshot = school_snapshot()
    for result in snapshot["results"]:
        if result["event_id"] == 4:
            result["event_name"] = "Gauteng North Championship"
            result["competition_type"] = LEAGUE  # Also supports already imported events.
    summary, _, rows = qualified_schools_report(snapshot)
    assert summary[0]["Qualified"] == "Yes"
    assert all(r["Completed league/championship meets"] == 4 for r in rows)
    snapshot["results"] = [r for r in snapshot["results"] if r["event_id"] != 3]
    summary, _, rows = qualified_schools_report(snapshot)
    assert summary[0]["Qualified"] == "No"
    assert all(r["Completed distinct qualifying events"] == 3 for r in rows)
    assert rows[0]["Missing to qualify"] == "1 more completed event"
    snapshot["results"] = [r for r in snapshot["results"] if r["event_id"] != 2]
    assert qualified_schools_report(snapshot)[0][0]["Qualified"] == "No"


def test_two_leagues_championship_and_interprovincial_qualify():
    snapshot = school_snapshot()
    for result in snapshot["results"]:
        if result["event_id"] == 3:
            result["event_name"] = "GN Championship"
            result["competition_type"] = CHAMPIONSHIP
    summary, _, rows = qualified_schools_report(snapshot)
    assert summary[0]["Qualified"] == "Yes"
    assert all(r["Completed distinct qualifying events"] == 4 for r in rows)


@pytest.mark.parametrize("types,qualified", [
    ([LEAGUE, LEAGUE, INTERPROVINCIAL, INTERPROVINCIAL], "Yes"),
    ([LEAGUE, INTERPROVINCIAL, INTERPROVINCIAL, INTERPROVINCIAL], "Yes"),
    ([LEAGUE] * 4, "No"),
    ([LEAGUE, LEAGUE, CHAMPIONSHIP], "No"),
])
def test_any_four_completed_meets_require_interprovincial_or_championship(types, qualified):
    snapshot = school_snapshot()
    snapshot["results"] = [r for r in snapshot["results"] if r["event_id"] <= len(types)]
    for result in snapshot["results"]:
        result["competition_type"] = types[result["event_id"] - 1]
        result["event_name"] = "GN Championship" if result["competition_type"] == CHAMPIONSHIP else f"Meet {result['event_id']}"
    assert qualified_schools_report(snapshot)[0][0]["Qualified"] == qualified


def test_unknown_school_visible_without_invented_score():
    summary, _, rows = qualified_schools_report(school_snapshot("Unknown School"))
    assert summary[0]["Total school points"] == ""
    assert summary[0]["Qualified"] == "No"
    assert len(rows) == 6


def test_rank_and_workbook_colours():
    snapshot = school_snapshot()
    other = school_snapshot("Second Hs", (15, 17, 19))
    for athlete in other["athletes"]:
        athlete["id"] += 100
        athlete["athlete_number"] += "0"
    for result in other["results"]:
        result["athlete_id"] += 100
        result["athlete_number"] += "0"
    snapshot["athletes"] += other["athletes"]
    snapshot["results"] += other["results"]
    summary, headers, rows = qualified_schools_report(snapshot)
    assert [r["School"] for r in summary] == ["Example Ps", "Second Hs"]
    assert [r["Rank"] for r in summary] == [1, 2]
    workbook = openpyxl.load_workbook(build_qualified_schools_xlsx(summary, headers, rows))
    assert workbook.sheetnames == ["Qualified Schools", "School Athlete Scores", "Primary Schools", "High Schools"]
    sheet = workbook["Qualified Schools"]
    assert sheet["B2"].fill.fgColor.rgb == "00C6EFCE"
    assert sheet["E2"].value == "Yes" and sheet["E2"].fill.fgColor.rgb == "00C6EFCE"
    assert sheet["E3"].value == "No" and sheet["E3"].fill.fgColor.rgb == "00FFC7CE"
    assert sheet.auto_filter.ref == sheet.dimensions


def test_school_type_tabs_rank_separately_and_summarize_missing_places():
    summary, headers, athletes = qualified_schools_report(school_snapshot("Example Hs", (19, 19, 19, 19)))
    assert summary[0]["Missing team members"] == "u/15; u/17"
    base = summary[0]
    summary.extend([
        {**base, "School": "Top Ps", "School type": "Primary", "Total school points": 12000,
         "Qualified": "Yes", "Qualification summary": "None", "Missing to qualify": "None"},
        {**base, "School": "Second Ps", "School type": "Primary", "Total school points": 10000},
        {**base, "School": "Tied Hs", "Total school points": base["Total school points"]},
    ])
    workbook = openpyxl.load_workbook(build_qualified_schools_xlsx(summary, headers, athletes))
    primary, high = workbook["Primary Schools"], workbook["High Schools"]
    assert list(next(high.values)) == [h.upper() for h in ["Rank", "School", "Total school points", "Qualified", "Full team", "Missing team members", "Missing Events", "Qualification requirements", "Missing to qualify detail"]]
    assert all(c.value == c.value.upper() for c in primary[1])
    assert primary["A2"].value == 1 and primary["A3"].value == 2
    assert high["A2"].value == high["A3"].value == 1
    assert high["E2"].value == "No"
    assert high["F2"].value == "u/15; u/17"
    assert "Missing mandatory athlete(s):" in high["I2"].value
    assert primary["B2"].fill.fgColor.rgb == primary["D2"].fill.fgColor.rgb == "00C6EFCE"
    assert high["D2"].fill.fgColor.rgb == "00FFC7CE"
    assert primary.column_dimensions["A"].width < primary.column_dimensions["C"].width
    assert high.column_dimensions["I"].width == 65
    assert high["I2"].alignment.wrap_text
    assert high.auto_filter.ref == high.dimensions
    for name in ("Qualified Schools", "Primary Schools", "High Schools"):
        assert {"FULL TEAM", "MISSING TEAM MEMBERS", "MISSING EVENTS", "QUALIFICATION REQUIREMENTS"}.issubset(h.upper() for h in next(workbook[name].values))


def test_compact_team_members_and_event_shortfalls():
    snapshot = school_snapshot(ages=(11, 15))
    snapshot["results"] = [r for r in snapshot["results"] if r["event_id"] <= 2]
    row = qualified_schools_report(snapshot)[0][0]
    assert row["Missing team members"] == "u/09; u/13; + 2"
    assert row["Missing Events"] == "1 x League; 1 x IP/Champs"
    # Missing attendance is summarized across the team, not added per athlete.
    assert "2 x IP" not in row["Missing Events"]


def test_school_summary_separates_athlete_shortfalls_from_detail():
    snapshot = school_snapshot()
    snapshot["results"] = [r for r in snapshot["results"] if r["athlete_id"] != 1 or r["event_id"] == 1]
    summary, _, _ = qualified_schools_report(snapshot)
    row = summary[0]
    assert row["Full team"] == "Yes"
    assert row["Missing team members"] == "None"
    assert "Under-9: interprovincial or GN Championship required" in row["Qualification requirements"]
    assert "Under-9: more completed" in row["Qualification requirements"]
    assert "Athlete 1" not in row["Qualification requirements"]
    assert "Athlete 1" in row["Missing to qualify"]
    assert "3 more completed events" in row["Missing to qualify"]


def test_school_event_summary_does_not_require_extra_league_when_only_ip_is_missing():
    snapshot = school_snapshot()
    snapshot["results"] = [r for r in snapshot["results"] if r["athlete_id"] != 1 or r["event_id"] != 4]
    row = qualified_schools_report(snapshot)[0][0]
    assert row["Full team"] == "Yes"
    assert row["Qualified"] == "No"
    assert row["Qualification requirements"] == "Under-9: interprovincial or GN Championship required"


@pytest.mark.parametrize("ip_layout", [False, True])
@pytest.mark.parametrize("extension", ["xlsx", "pdf"])
def test_championship_import(ip_layout, extension):
    fixture = excel_fixture if extension == "xlsx" else pdf_fixture
    event = parse_season_results(fixture(ip_layout, title="GAUTENG NORTH CHAMPIONSHIP"), f"championship.{extension}")
    assert event.competition_type == CHAMPIONSHIP
    assert event.results and event.awards


def test_high_school_missing_group_still_counts_three_extras():
    summary, _, details = qualified_schools_report(school_snapshot("Example Hs", (15, 17, 15, 17, 15, "SPECIAL NEEDS")))
    assert summary[0]["Mandatory places filled"] == "2/3"
    assert summary[0]["Extra places filled"] == "3/3"
    assert sum(r["Selected"] == "Yes" for r in details) == 5


@pytest.mark.parametrize("ages,mandatory_count,extra_count,total", [
    ((17, 19, 17, 19, 17, 19), 2, 3, 11250),
    ((19, 19, 19, 19, 19, 19), 1, 3, 9200),
    ((19, 19, 19, 19), 1, 3, 9200),
    ((19, 19), 1, 1, 4800),
])
def test_high_school_partial_team_counts_only_available_places(ages, mandatory_count, extra_count, total):
    summary, headers, athletes = qualified_schools_report(school_snapshot("Example Hs", ages))
    row = summary[0]
    assert row["Qualified"] == "No"
    assert row["Mandatory places filled"] == f"{mandatory_count}/3"
    assert row["Extra places filled"] == f"{extra_count}/3"
    assert row["Total school points"] == total
    selected = [a for a in athletes if a["Selected"] == "Yes"]
    assert len(selected) == mandatory_count + extra_count
    assert len({a["Athlete number"] for a in selected}) == len(selected)
    assert {a["Team place"] for a in selected if a["Team place"].startswith("Extra")} == {f"Extra {i}" for i in range(1, extra_count + 1)}
    assert sum(a["School contribution points"] for a in athletes) == total
    workbook = openpyxl.load_workbook(build_qualified_schools_xlsx(summary, headers, athletes))
    assert workbook["Qualified Schools"]["D2"].value == total


def test_capped_score_used_for_selection_not_raw_special_score():
    snapshot = school_snapshot(ages=(9, 11, 13, 15, 11, 13, "SPECIAL NEEDS"))
    for result in snapshot["results"]:
        if result["athlete_id"] in {5, 6}:
            result["total_points"] = "2100"
        elif result["athlete_id"] == 7:
            result["total_points"] = "4000"
    _, _, details = qualified_schools_report(snapshot)
    special = next(r for r in details if r["Athlete number"] == "007")
    assert special["Selected"] == "No"
    assert special["Best total points"] == 4000


def test_latest_school_and_category_use_event_date_not_import_order():
    snapshot = school_snapshot()
    for result in snapshot["results"]:
        if result["athlete_id"] == 1 and result["event_id"] == 4:
            result["school"] = "New Hs"
            result["category"] = "U/15 GIRLS"
    snapshot["results"].reverse()
    summary, _, details = qualified_schools_report(snapshot)
    athlete = next(r for r in details if r["Athlete number"] == "001")
    assert athlete["School"] == "New Hs" and athlete["Age group"] == "U/15 GIRLS"
    assert sum(r["Athlete number"] == "001" for r in details) == 1
    original = next(r for r in summary if r["School"] == "Example Ps")
    assert "Under-9" in original["Missing to qualify"]


def test_championship_store_roundtrip_and_duplicate(tmp_path):
    from database.season_db import init_season_db
    from database.season_repository import store_event, find_event, season_snapshot, EventAlreadyExists
    db = tmp_path / "season.sqlite"
    init_season_db(db)
    event = parse_season_results(excel_fixture(title="GN CHAMPIONSHIPS"), "champ.xlsx")
    store_event(db, event)
    assert find_event(db, event)["competition_type"] == CHAMPIONSHIP
    with pytest.raises(EventAlreadyExists):
        store_event(db, event)
    store_event(db, event, overwrite=True)
    assert len(season_snapshot(db)["events"]) == 1
