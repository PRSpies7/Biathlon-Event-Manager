"""Small synthetic fixtures derived from the four supplied layouts, no personal data."""
from dataclasses import replace
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
import sqlite3

import openpyxl
import pytest

from database.db import get_conn, init_db
from database.repository import create_event, delete_all_events, get_events
from database.season_db import init_season_db, season_database_path
from database.season_repository import EventAlreadyExists, find_event, reset_season_database, season_snapshot, set_affiliation, store_event
from exporters.season_results import build_season_results_xlsx
from parsers.season_results import parse_season_results
from parsers.season_results.common import clean_name, recognize_metadata
from parsers.season_results.models import Award, INTERPROVINCIAL, LEAGUE, NormalizedEvent, Result
from services.season_results_service import athlete_rows, import_event, import_summary

LEAGUE_HEADER = [None, "#", "Athlete", "Run Time", "RP", "Swim Time", "SP", "BP", "Total Points"]
IP_HEADER = [None, "#", "Athlete", "Province", "Team", "Running Time", "Running Points", "Swimming Time", "Swimming Points", "Bonus Points", "Total Points"]


def sample_event():
    return NormalizedEvent("GN League 1", date(2026, 8, 25), LEAGUE, "results.xlsx", [
        Result("00101", "Alex Example", "U/11 GIRLS", "1", "01:20.10", "1040.00", "00:40.50", "1020.00", "0.00", "2060.00", school="Example School", affiliated=True, team="GN"),
        Result("102", "Robin Example", "U/11 GIRLS", "2", "DNF", "DNF", "DNF", "DNF", "DNF", "DNF", status="DNF", team="GN"),
    ], [Award("00101", "Alex Example", kind, i + 1) for i, kind in enumerate(("Runner", "Swimmer", "Overall Athlete"))])


@pytest.fixture
def db(tmp_path):
    path = tmp_path / "season.sqlite"
    init_season_db(path)
    return path


def excel_fixture(interprovincial=False, title=None):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([title or ("GN INTER-PROVINCIAL" if interprovincial else "GN LEAGUE 1")])
    sheet.append([datetime(2026, 8, 25) if interprovincial else "Date: 2026-08-25"])
    header = IP_HEADER if interprovincial else LEAGUE_HEADER
    names = ["Alex Example (AFL*) (Example School)" if not interprovincial else "Alex Example (AFL*)",
             "Robin Example (TO)", "Sam Example (AFL)"]
    rows = [[i + 1, "00101" if i == 0 else str(102 + i - 1), name,
             *(["GAUTENG NORTH", "GN Team A"] if interprovincial else []),
             time(0, 1, 20, 100000), 1040, time(0, 0, 40, 500000), 1020, 0, "2060.00\u00a0PB"] for i, name in enumerate(names)]
    for kind in ("Runners", "Swimmers", "Athletes"):
        sheet.append(header)
        sheet.append([f"Top 3 {kind}"])
        for row in rows:
            sheet.append(row)
    sheet.append(header)
    sheet.append(["U/11 GIRLS"])
    for row in rows:
        sheet.append(row)
    sheet.append([4, "104", "Finisher Example", *(["GN", "GN"] if interprovincial else []), *(["DNF"] * 6)])
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def pdf_fixture(interprovincial=False, title=None):
    """Build tiny text PDFs using standard PDF primitives, without extra dependencies.

    League tables have vertical rules. IP tables only rule the headers, change
    column widths between sections, and split a name/team across two pages.
    """
    pages, commands = [], []
    height = 850

    def label(x, y, value):
        escaped = str(value or "").replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(f"BT /F1 9 Tf {x} {height-y} Td ({escaped}) Tj ET")

    def line(x, y, x2, y2):
        commands.append(f"{x} {height-y} m {x2} {height-y2} l S")

    def table(y, heading, number, name, award=False, split=False):
        columns = ([20, 48, 96, 213, 271, 347, 423, 499, 575, 638, 725] if interprovincial
                   else [20, 48, 96, 271, 347, 423, 499, 575, 650, 725])
        header = ([None, "#", "Athlete", "Team", "Running Time", "Running Points", "Swimming Time", "Swimming Points", "Bonus Points", "Total Points"]
                  if interprovincial else LEAGUE_HEADER)
        # Exercise changed header-derived widths, not hard-coded x coordinates.
        if award and interprovincial:
            columns[3] += 8
        for i, h in enumerate(header):
            label(columns[i] + 3, y + 14, h)
        label(23, y + 37, heading)
        for boundary in (y, y + 23, y + 46):
            line(20, boundary, 725, boundary)
        for x in columns:
            line(x, y, x, y + 23)
            if not interprovincial:
                line(x, y + 46, x, y + 100)
        if not interprovincial:
            line(columns[0], y + 23, columns[0], y + 46)
            line(columns[-1], y + 23, columns[-1], y + 46)
        row = [1, number, name, *(["GN Team"] if interprovincial else []), "01:20.100", "1040.00", "00:40.500", "1020.00", "0.00", "2060.00 PB"]
        for i, val in enumerate(row):
            label(columns[i] + 3, y + 63, val)
        if split:
            pages.append("\n".join(commands))
            commands.clear()
            label(columns[2] + 3, 20, "Example (AFL*)")
            label(columns[3] + 3, 20, "A")
            line(20, 32, 725, 32)
        else:
            label(columns[2] + 3, y + 77, "Example (AFL*) (Example School)" if not interprovincial else "Example (AFL*)")
            if interprovincial:
                label(columns[3] + 3, y + 77, "A")
            line(20, y + 100, 725, y + 100)

    label(20, 24, title or ("GN INTER-PROVINCIAL" if interprovincial else "GN LEAGUE 1"))
    label(20, 42, "Date: 2026-08-25")
    for i, kind in enumerate(("Runners", "Swimmers", "Athletes")):
        table(60 + i * 130, f"Top 3 {kind}", "101", "Alex", award=True)
    table(690 if interprovincial else 450, "SPECIAL NEEDS FEMALE", "101", "Alex", split=interprovincial)
    pages.append("\n".join(commands))
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    page_ids = []
    for content in pages:
        page_id = len(objects) + 1
        page_ids.append(page_id)
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 750 {height}] /Resources << /Font << /F1 3 0 R >> >> /Contents {page_id+1} 0 R >>".encode())
        stream = content.encode("ascii")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    objects[1] = f"<< /Type /Pages /Count {len(pages)} /Kids [{' '.join(f'{i} 0 R' for i in page_ids)}] >>".encode()
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    start = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer << /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF".encode())
    return BytesIO(output)


@pytest.mark.parametrize("interprovincial", [False, True])
def test_excel_import(interprovincial):
    event = parse_season_results(excel_fixture(interprovincial), "wrong-date-2020-01-01.xlsx")
    event.check_structure()
    assert event.event_date == date(2026, 8, 25)
    assert event.competition_type == (INTERPROVINCIAL if interprovincial else LEAGUE)
    assert len(event.results) == 4 and len(event.awards) == 9
    assert event.results[0].athlete_number == "00101"
    assert event.results[0].athlete_name == "Alex Example"
    assert event.results[0].total_points == "2060.00"
    assert event.results[0].run_time == "01:20.10"
    assert event.results[0].affiliated and event.results[2].affiliated
    assert not event.results[1].affiliated
    assert event.results[-1].status == "DNF"
    assert event.results[-1].total_points == "DNF"
    assert event.results[0].team == ("GN Team A" if interprovincial else "")
    assert event.results[0].school == ("" if interprovincial else "Example School")


@pytest.mark.parametrize("interprovincial", [False, True])
def test_pdf_import_wrapped_names_columns_and_page_continuation(interprovincial):
    event = parse_season_results(pdf_fixture(interprovincial), "incorrect-2025-08-25.pdf")
    event.check_structure()
    assert event.event_date == date(2026, 8, 25)
    assert len(event.results) == 1 and len(event.awards) == 3
    result = event.results[0]
    assert result.athlete_name == "Alex Example"
    assert result.affiliated
    assert result.total_points == "2060.00"
    assert result.category == "SPECIAL NEEDS FEMALE"
    assert result.team == ("GN Team A" if interprovincial else "")
    assert result.school == ("" if interprovincial else "Example School")


@pytest.mark.parametrize("marker,expected", [("(AFL)", True), ("(AFL*)", True), ("( afl * )", True), ("(TO)", False), ("PB", False), ("", False)])
def test_affiliation_and_name_annotations(marker, expected):
    name, affiliated, school, annotations = clean_name(f"Alex Example {marker} (Example School)", True)
    assert name == "Alex Example" and affiliated is expected and school == "Example School"
    assert annotations == ("TO" if marker == "(TO)" else "")


@pytest.mark.parametrize("kind", ["Runner", "Swimmer", "Overall Athlete"])
def test_structured_awards_and_placing(kind):
    event = parse_season_results(excel_fixture(), "results.xlsx")
    awards = [a for a in event.awards if a.award_type == kind]
    assert [(a.athlete_number, a.placing) for a in awards] == [("00101", 1), ("102", 2), ("103", 3)]


@pytest.mark.parametrize("value", ["Date: 2026-08-25", "25-08-2026", "25/08/2026", "2026/08/25"])
def test_date_recognition(value):
    assert recognize_metadata(["GN LEAGUE 1", value], "wrong-2025.pdf").event_date == date(2026, 8, 25)


def test_national_style_title_requires_confirmation_but_uses_ip_structure():
    event = parse_season_results(excel_fixture(True, "SA BIATHLON CHAMPIONSHIPS 2026"), "Inter-Provincial Example.xlsx")
    assert event.competition_type == "" and event.notes
    with pytest.raises(ValueError, match="supported competition"):
        event.check_structure()
    replace(event, name="GN Interprovincial", competition_type=INTERPROVINCIAL).check_structure()


def test_duplicate_uses_identity_date_type_not_filename(db):
    event = sample_event()
    event_id = store_event(db, event)
    duplicate = replace(event, name="  gn LEAGUE-1 ", source_filename="renamed.pdf")
    assert find_event(db, duplicate)["id"] == event_id
    assert import_summary(db, duplicate)["Status"] == "Already Imported"
    with pytest.raises(EventAlreadyExists):
        store_event(db, duplicate)
    assert find_event(db, replace(event, event_date=date(2026, 9, 1))) is None
    assert find_event(db, replace(event, competition_type=INTERPROVINCIAL)) is None


def test_complete_overwrite_replaces_results_and_awards(db):
    event = sample_event()
    event_id = store_event(db, event)
    corrected = replace(event, results=[replace(event.results[0], total_points="1234.56", school="", affiliated=False)],
                        awards=[Award("00101", "Alex", "Swimmer", 1)], source_filename="corrected.pdf")
    assert store_event(db, corrected, overwrite=True) == event_id
    saved = season_snapshot(db)
    assert len(saved["events"]) == 1 and saved["events"][0]["result_count"] == 1
    assert len(saved["results"]) == len(saved["awards"]) == 1
    assert saved["results"][0]["total_points"] == "1234.56"
    assert saved["awards"][0]["award_type"] == "Swimmer" and saved["awards"][0]["placing"] == 1
    alex = next(a for a in saved["athletes"] if a["athlete_number"] == "00101")
    assert alex["affiliated"] and alex["school"] == "Example School"
    assert next(a for a in saved["athletes"] if a["athlete_number"] == "102")["events_attended"] == 0


def test_overwrite_failure_rolls_back_all_changes(db):
    event = sample_event()
    store_event(db, event)
    before = season_snapshot(db)
    with get_conn(db) as conn:
        conn.execute("CREATE TRIGGER fail_result BEFORE INSERT ON season_results BEGIN SELECT RAISE(ABORT, 'test failure'); END")
    with pytest.raises(sqlite3.IntegrityError):
        store_event(db, replace(event, results=[replace(event.results[0], athlete_name="Changed")], awards=[]), overwrite=True)
    assert season_snapshot(db) == before


def test_identity_by_number_and_affiliation_cannot_regress(db):
    event = sample_event()
    store_event(db, event)
    later = replace(event, name="GN IP", competition_type=INTERPROVINCIAL, results=[
        replace(event.results[0], athlete_name="Alex Corrected", school="", team="GN Team A", affiliated=False),
        replace(event.results[1], athlete_name="Alex Example", affiliated=True),
    ], awards=[])
    store_event(db, later)
    athletes = {a["athlete_number"]: a for a in season_snapshot(db)["athletes"]}
    assert len(athletes) == 2
    assert athletes["00101"]["athlete_name"] == "Alex Corrected"
    assert athletes["00101"]["school"] == "Example School"
    assert athletes["00101"]["team"] == "GN Team A"
    assert athletes["00101"]["affiliated"] == athletes["102"]["affiliated"] == 1
    assert athletes["00101"]["events_attended"] == 2


def test_manual_affiliation_can_correct_either_direction(db):
    store_event(db, sample_event())
    athlete_id = season_snapshot(db)["athletes"][0]["id"]
    set_affiliation(db, athlete_id, False)
    assert season_snapshot(db)["athletes"][0]["affiliated"] == 0
    set_affiliation(db, athlete_id, True)
    assert season_snapshot(db)["athletes"][0]["affiliated"] == 1


def test_skip_has_no_side_effects_and_batch_duplicates_are_guarded(db):
    event = sample_event()
    assert import_event(db, event, "Skip") is None
    assert season_snapshot(db)["events"] == []
    import_event(db, event)
    before = season_snapshot(db)
    assert import_event(db, event, "Skip") is None
    assert season_snapshot(db) == before
    with pytest.raises(EventAlreadyExists):
        import_event(db, event)


def test_award_marker_promotes_affiliation(db):
    event = sample_event()
    event.results[0].affiliated = False
    event.awards[0].affiliated = True
    store_event(db, event)
    assert season_snapshot(db)["athletes"][0]["affiliated"] == 1


def test_export_filterable_athletes_structured_awards_and_published_values(db):
    event = sample_event()
    event.results[0].athlete_name = "=Literal name"
    store_event(db, event)
    snapshot = season_snapshot(db)
    workbook = openpyxl.load_workbook(build_season_results_xlsx(snapshot))
    assert workbook.sheetnames == ["Athletes", "Event Results", "Imported Events", "Event Awards"]
    sheet = workbook["Athletes"]
    assert sheet.max_row == 3 and sheet.freeze_panes == "C2" and sheet.auto_filter.ref == "A1:M3"
    assert sheet["A2"].value == "00101" and sheet["A2"].data_type == "s"
    assert sheet["B2"].value == "=Literal name" and sheet["B2"].data_type == "s"
    assert {sheet["G2"].value, sheet["G3"].value} == {"Affiliated", "Not Affiliated"}
    assert sheet["I2"].value == athlete_rows(snapshot)[0]["Event awards"]
    assert "GN League 1 - 1st Top Runner" in sheet["I2"].value
    assert "3rd Top Athlete" in sheet["I2"].value
    result_sheet = workbook["Event Results"]
    assert result_sheet["N2"].value == 2060 and result_sheet["N3"].value == "DNF"
    assert workbook["Event Awards"].max_row == 4
    assert isinstance(workbook["Imported Events"]["C2"].value, datetime)
    workbook.close()


def test_missing_number_and_duplicate_numbers_are_not_fuzzy_matched(db):
    event = sample_event()
    with pytest.raises(ValueError, match="missing"):
        store_event(db, replace(event, results=[replace(event.results[0], athlete_number="")]))
    with pytest.raises(ValueError, match="Repeated"):
        store_event(db, replace(event, results=[event.results[0], event.results[0]]))
    assert not season_snapshot(db)["events"]


def test_workflow_reset_cannot_delete_season_results(db):
    init_db(db)
    store_event(db, sample_event())
    create_event(str(db), "Workflow", "Host", "2026-08-25", "SCM", 8)
    before = season_snapshot(db)
    delete_all_events(str(db))
    assert not get_events(str(db))
    assert season_snapshot(db) == before
    init_season_db(db)
    assert season_snapshot(db) == before


def test_persistent_database_configuration(tmp_path):
    assert season_database_path(tmp_path, {}) == tmp_path / "data" / "season_results.sqlite"
    assert season_database_path(tmp_path, {"DATABASE_URL": "sqlite:///custom.sqlite"}) == tmp_path / "custom.sqlite"
    assert season_database_path(tmp_path, {"SEASON_DATABASE_PATH": "season-2027.sqlite"}) == tmp_path / "season-2027.sqlite"
    for url in ("postgresql://example.invalid/db", "sqlite:///:memory:", "sqlite:///"):
        with pytest.raises(ValueError):
            season_database_path(tmp_path, {"DATABASE_URL": url})


def test_missing_date_never_uses_misleading_filename():
    event = recognize_metadata(["GN LEAGUE 1"], "2025-08-25.pdf")
    assert event.event_date is None and event.notes


def test_reset_clears_entire_season_preserves_workflow_and_can_reimport(db):
    init_db(db)
    create_event(str(db), "Workflow", "Host", "2026-08-25", "SCM", 8)
    workflow = get_events(str(db))
    store_event(db, sample_event())
    reset_season_database(db)
    assert all(not rows for rows in season_snapshot(db).values())
    assert get_events(str(db)) == workflow
    store_event(db, sample_event())
    assert len(season_snapshot(db)["results"]) == 2


def test_reset_failure_rolls_back_entire_season(db):
    store_event(db, sample_event())
    before = season_snapshot(db)
    with get_conn(db) as conn:
        conn.execute("CREATE TRIGGER fail_reset BEFORE DELETE ON season_athletes BEGIN SELECT RAISE(ABORT, 'test failure'); END")
    with pytest.raises(sqlite3.IntegrityError):
        reset_season_database(db)
    assert season_snapshot(db) == before
