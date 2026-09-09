from copy import deepcopy, copy
from dataclasses import replace
from io import BytesIO

import openpyxl
import pytest

from database.season_db import init_season_db
from database.season_repository import store_event, season_snapshot, replace_record_benchmarks, reset_season_database
from exporters.season_records import updated_records_xlsx
from exporters.season_results import build_top_athletes_xlsx, build_qualified_athletes_xlsx, with_insights, build_season_results_xlsx
from parsers.season_records import parse_record_benchmarks, category_key
from parsers.season_results.models import INTERPROVINCIAL
from services.season_insights import insight_sections
from services.season_reports import qualified_report_rows, score_qualification, filtered_snapshot
from services.season_scope import is_gn_team
from test_season_results import sample_event
from test_season_reports import add_events
from test_school_reports import school_snapshot
from services.school_reports import qualified_schools_report


@pytest.fixture
def db(tmp_path):
    path = tmp_path / "season.sqlite"
    init_season_db(path)
    return path


@pytest.mark.parametrize("team,province,expected", [("GN", "", True), ("Team GN", "", True), ("GN Team A", "", True),
    ("GAUTENG NORTH", "", True), ("", "Gauteng North", True), ("NW", "GN", False), ("LIM", "", False),
    ("SFS", "", False), ("", "", False), ("GNARLY", "", False)])
def test_exact_gn_scope(team, province, expected):
    assert is_gn_team(team, province) is expected


def test_import_excludes_non_gn_athletes_and_awards(db):
    base = sample_event()
    event = replace(base, competition_type=INTERPROVINCIAL,
                    results=[base.results[0], replace(base.results[1], team="NW")])
    store_event(db, event)
    snapshot = season_snapshot(db)
    assert [a["athlete_number"] for a in snapshot["athletes"]] == ["00101"]
    assert len(snapshot["awards"]) == 3
    before = deepcopy(snapshot)
    assert season_snapshot(db) == before  # Reads do not clean historical imports.


@pytest.mark.parametrize("category,score,missing", [("U/11 GIRLS", "1889.99", True), ("U/11 GIRLS", "1890", False),
    ("JNR WOMEN", "1789.99", True), ("JNR WOMEN", "1790", False), ("SENIORS MEN", "1790", False),
    ("MASTERS 70+ MEN", "1790", False)])
def test_score_thresholds(category, score, missing):
    r = {"category": category, "total_points": score, "run_time": "01:20", "swim_time": "00:40", "status": ""}
    assert bool(score_qualification([r], category)) == missing
    r["status"] = "DNF"
    assert score_qualification([r], category)


def test_school_score_requirement():
    snapshot = school_snapshot()
    for r in snapshot["results"]:
        if r["athlete_id"] == 1:
            r["total_points"] = "1889.99"
    summary = qualified_schools_report(snapshot)[0][0]
    assert summary["Qualified"] == "No"
    assert "1890" in summary["Missing to qualify"]


def test_report_sheets_colours_and_insights(db):
    snapshot = add_events(db)
    rows = qualified_report_rows(snapshot)
    workbook = openpyxl.load_workbook(build_qualified_athletes_xlsx(snapshot, rows))
    assert workbook.sheetnames == ["Qualified Athletes", "Insights"]
    sheet = workbook.active
    assert sheet["C2"].fill.fgColor.rgb == sheet["D2"].fill.fgColor.rgb == "00C6EFCE"
    for r in snapshot["results"]:
        r["total_points"] = "1800"
    rows = qualified_report_rows(snapshot)
    workbook = openpyxl.load_workbook(build_qualified_athletes_xlsx(snapshot, rows))
    assert workbook.active["C2"].fill.fgColor.rgb == workbook.active["D2"].fill.fgColor.rgb == "00FFC7CE"
    top = openpyxl.load_workbook(build_top_athletes_xlsx(snapshot))
    assert top.sheetnames == ["Top Athletes", "Insights", "Event Details"]
    assert {"Run time", "Swim time", "Running points", "Swimming points"}.issubset(next(top["Event Details"].values))
    assert len(top["Insights"]._charts) == 2
    followup = insight_sections(snapshot, affiliation_only=True)[0][2]
    assert len(followup) == 1 and followup[0]["Events attended"] == 4
    affiliated = filtered_snapshot(snapshot, set())
    output = openpyxl.load_workbook(with_insights(build_season_results_xlsx(affiliated, include_points=False), snapshot, affiliation_only=True))
    assert output.sheetnames[1] == "Insights"
    assert all("points" not in str(cell.value).casefold() for sheet in output for row in sheet for cell in row)


def template():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["B2"] = "GAUTENG NORTH BIATHLON RECORDS - 2025"
    headers = ["AGE GROUP", "NAME OF RECORD HOLDER", "OLD RECORD", "RECORD POINTS", "RECORD TYPE", "SEASON", "90%", "AGE GROUP SA RECORD", "MEET NAME", "RUN TIME", "SWIM TIME"]
    for i, h in enumerate(headers, 2):
        sheet.cell(3, i, h)
    for row, category in [(4, "GIRLS U/11"), (5, "BOYS U/11")]:
        for col, value in {2: category, 3: "Previous Holder", 4: 2200, 5: 2000, 6: "(GN)", 7: 2025, 8: f"=D{row}*0.9", 9: 2300, 10: "Previous meet", 11: "01:20", 12: "00:40"}.items():
            sheet.cell(row, col, value)
    sheet.column_dimensions["C"].width = 35
    sheet["C4"].font = openpyxl.styles.Font(bold=True)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_supplied_record_layout_persistence_and_updated_download(db):
    payload = template()
    references = parse_record_benchmarks(BytesIO(payload))
    assert len(references) == 2
    assert references[0]["athlete_number"] == ""
    assert category_key("GIRLS U/11") == category_key("U/11 GIRLS")
    assert references[0]["total_points"] == "2000.00"
    replace_record_benchmarks(db, references, "records.xlsx", payload)
    snapshot = add_events(db)
    out = openpyxl.load_workbook(updated_records_xlsx(snapshot))
    source = openpyxl.load_workbook(BytesIO(payload))
    sheet = out.worksheets[0]
    assert sheet["C4"].value == "Alex Example" and sheet["E4"].value == 2060
    assert sheet["D4"].value == 2000  # RECORD POINTS, not OLD RECORD, was the baseline.
    assert sheet["H4"].value == "=D4*0.9" and sheet["I4"].value == 2300
    assert copy(sheet["C4"].font) == copy(source.active["C4"].font)
    assert sheet.column_dimensions["C"].width == 35
    assert [c.value for c in sheet[5]] == [c.value for c in source.active[5]]
    assert out["New Records"].max_row == 2
    assert out["New Records"]["G2"].value == "Event 1"
    assert out["New Records"]["H2"].value.year == 2026
    assert season_snapshot(db)["record_templates"][0]["content"] == payload
    assert season_snapshot(db)["records"] == snapshot["records"]  # Download doesn't alter saved references.
    reset_season_database(db)
    assert not season_snapshot(db)["record_templates"]


def test_equal_records_do_not_replace_holder(db):
    payload = template()
    replace_record_benchmarks(db, [], "records.xlsx", payload)
    event = sample_event()
    store_event(db, replace(event, results=[replace(event.results[0], total_points="2000")]))
    output = openpyxl.load_workbook(updated_records_xlsx(season_snapshot(db)))
    assert output.active["C4"].value == "Previous Holder"
    assert output["New Records"].max_row == 1


def test_under_11_girls_record_replaced_by_u11_female_result(db):
    workbook = openpyxl.load_workbook(BytesIO(template()))
    workbook.active["B4"] = "Under 11 Girls"
    payload = BytesIO()
    workbook.save(payload)
    replace_record_benchmarks(db, [], "records.xlsx", payload.getvalue())
    event = sample_event()
    store_event(db, replace(event, results=[replace(event.results[0], category="U/11 Female", total_points="2100")]))
    updated = openpyxl.load_workbook(updated_records_xlsx(season_snapshot(db)))
    assert updated.active["B4"].value == "Under 11 Girls"
    assert updated.active["C4"].value == "Alex Example"
    assert updated.active["D4"].value == 2000  # Previous Record Points becomes Old Record.
    assert updated.active["E4"].value == 2100  # New score replaces Record Points.
