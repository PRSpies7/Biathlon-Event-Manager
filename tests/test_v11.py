from pathlib import Path

from openpyxl import load_workbook

from database.db import init_db
from database.repository import clear_results, create_event, get_athletes, replace_athletes, update_manual_time, update_run_times, update_swim_times
from exporters.swim_timekeeper import build_swim_timekeeper_xlsx
from validation.validators import normalize_time


def _db(tmp_path):
    p = tmp_path / "test.sqlite"
    init_db(p)
    return str(p)


def test_time_normalization_variants():
    expected = "04:12.50"
    for value in ["04:12.50", "4:12.50", "04;12;50", "04,12,50", "04.12.50", "04 12 50", "04:12:50"]:
        assert normalize_time(value) == expected
    assert normalize_time("") is None
    assert normalize_time("0.00") is None
    assert normalize_time("04:99.50") is None


def test_manual_time_survives_reimport_and_reset(tmp_path):
    db = _db(tmp_path)
    event = create_event(db, "Test", "GN", "2026-08-14", "LCM", 6)
    replace_athletes(db, event, [{
        "sort_order": 1, "athlete_number": "1011", "athlete_name": "John Doe", "group_name": "U/19 MEN",
        "running_heat": 1, "running_lane": 1, "swimming_heat": 1, "swimming_lane": 1,
    }])
    update_run_times(db, event, {"1011": "04:12.50"})
    update_swim_times(db, event, {"1011": "01:05.20"})
    update_manual_time(db, event, "1011", "run_time", "04:11.90")
    update_run_times(db, event, {"1011": "04:20.00"})
    a = get_athletes(db, event)[0]
    assert a["run_time"] == "04:11.90"
    assert a["run_time_imported"] == "04:20.00"
    clear_results(db, event)
    a = get_athletes(db, event)[0]
    assert a["run_time"] is None and a["swim_time"] is None
    assert a["run_position"] is None
    assert a["running_heat"] == 1


def test_swim_workbook_has_lane_tabs_and_print_headers():
    athletes = [
        {"sort_order": 1, "athlete_number": "1011", "athlete_name": "John Doe", "group_name": "U/19 MEN", "swimming_heat": 1, "swimming_lane": 2},
        {"sort_order": 2, "athlete_number": "1012", "athlete_name": "Jane Doe", "group_name": "U/19 WOMEN", "swimming_heat": 2, "swimming_lane": 2},
    ]
    data = build_swim_timekeeper_xlsx(athletes, "Test Event", 6)
    path = Path("/tmp/test_swim_lanes.xlsx")
    path.write_bytes(data.getvalue())
    wb = load_workbook(path)
    assert wb.sheetnames == ["Lane1", "Lane2", "Lane3", "Lane4", "Lane5", "Lane6"]
    ws = wb["Lane2"]
    assert ws["A1"].value == "Test Event"
    assert ws["A2"].value == "Swim Timekeeper - Lane 2"
    assert ws["A4"].value.startswith("Timekeeper Name:")
    assert ws["B8"].value == "1011"
    assert ws.print_title_rows == "$1:$7"
    assert ws.page_setup.orientation == "portrait"
    assert ws.page_setup.fitToWidth == 1
