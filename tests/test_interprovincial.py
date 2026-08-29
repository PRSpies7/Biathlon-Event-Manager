from __future__ import annotations

import sqlite3
import tempfile
from pathlib import Path
import unittest

import openpyxl

from database.db import init_db
from database.repository import create_event, get_athletes, get_event, replace_athletes
from exporters.athlete_event_mapping import build_printable_athlete_mapping_xlsx
from exporters.swim_timekeeper import build_swim_timekeeper_xlsx
from exporters.timedrops_json import generate_timedrops_json
from parsers.master_entries import parse_master_entries


class InterprovincialPersistenceTests(unittest.TestCase):
    def test_repeated_header_workbook_preserves_province_and_assignments(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "entries.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["Running Heats"])
            sheet.append(["Lane", "Heat Number", "Group Name", "Athlete", "Athlete Number", "Province", "Time 1", "Time 2"])
            sheet.append([2, 3, "U/11 GIRLS (400m)", "Runner One", 101, "GN"])
            sheet.append([3, 4, "U/15 GIRLS (800m)", "Runner Two", 102, "LIM"])
            sheet.append(["Swimming Heats"])
            sheet.append(["Lane", "Heat Number", "Group Name", "Athlete", "Athlete Number", "Province", "Time 1", "Time 2"])
            sheet.append([5, 7, "U/11 GIRLS (50m)", "Runner One", 101, "GN"])
            sheet.append([6, 8, "U/15 GIRLS (100m)", "Runner Two", 102, "LIM"])
            workbook.save(path)

            parsed = parse_master_entries(path)

        self.assertEqual(parsed["source_type"], "xlsx_interprovincial")
        athletes = {row["athlete_number"]: row for row in parsed["athletes"]}
        self.assertEqual(athletes["101"]["province"], "GN")
        self.assertEqual(athletes["101"]["group_name"], "U/11 GIRLS")
        self.assertEqual((athletes["101"]["running_heat"], athletes["101"]["swimming_heat"]), (3, 7))
        self.assertEqual(athletes["102"]["province"], "LIM")
        self.assertEqual(athletes["102"]["group_name"], "U/15 GIRLS")

        mapping = openpyxl.load_workbook(build_printable_athlete_mapping_xlsx(parsed["athletes"]))
        mapping_groups = {mapping.active.cell(row, 3).value for row in range(2, mapping.active.max_row + 1)}
        self.assertEqual(mapping_groups, {"U/11 GIRLS", "U/15 GIRLS"})
        mapping.close()

        lane_sheets = openpyxl.load_workbook(build_swim_timekeeper_xlsx(parsed["athletes"], "IP", 8))
        exported_groups = {
            sheet.cell(row, 4).value
            for sheet in lane_sheets.worksheets
            for row in range(8, sheet.max_row + 1)
            if sheet.cell(row, 4).value
        }
        self.assertEqual(exported_groups, {"U/11 GIRLS", "U/15 GIRLS"})
        lane_sheets.close()

    def test_clean_groups_still_use_existing_json_age_group_distance_rules(self):
        athletes = [
            {"athlete_number": "101", "athlete_name": "One", "group_name": "U/11 GIRLS", "province": "GN", "swimming_heat": 7, "swimming_lane": 1},
            {"athlete_number": "102", "athlete_name": "Two", "group_name": "U/15 GIRLS", "province": "LIM", "swimming_heat": 8, "swimming_lane": 1},
        ]
        output = generate_timedrops_json(
            {}, athletes, meet_name="IP", host_team="Host", start_date="2026-08-25",
            course="LCM", pool_lanes=8, meet_type="Interprovincial",
        )
        races = output["meetSessions"][0]["sessionRaces"]
        self.assertEqual([(race["raceNumber"], race["raceEventNumber"]) for race in races], [(7, "2"), (8, "3")])
        self.assertEqual([event["eventDistance"] for event in output["meetEvents"]], [25, 50, 100])
        self.assertEqual(output["meetSwimmers"][0]["swimmerTeamId"], "1")
        self.assertEqual(races[0]["raceLanes"][0]["laneTeamId"], "1")

    def test_meet_type_and_province_round_trip_and_old_db_migration(self):
        with tempfile.TemporaryDirectory() as tmp:
            db_path = str(Path(tmp) / "events.sqlite")
            init_db(db_path)
            event_id = create_event(db_path, "IP", "Host", "2026-08-25", "LCM", 8, "Interprovincial")
            replace_athletes(db_path, event_id, [{
                "sort_order": 1,
                "athlete_number": "101",
                "athlete_name": "Runner One",
                "group_name": "U/08 GIRLS (400m)",
                "province": "GN",
                "running_heat": 3,
                "running_lane": 2,
                "swimming_heat": 7,
                "swimming_lane": 5,
            }])
            self.assertEqual(get_event(db_path, event_id)["meet_type"], "Interprovincial")
            self.assertEqual(get_athletes(db_path, event_id)[0]["province"], "GN")

            old_db = str(Path(tmp) / "old.sqlite")
            conn = sqlite3.connect(old_db)
            conn.executescript("""
                CREATE TABLE events (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, host_team TEXT NOT NULL,
                    start_date TEXT NOT NULL, course TEXT NOT NULL DEFAULT 'LCM',
                    pool_lanes INTEGER NOT NULL DEFAULT 8, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
                );
            """)
            conn.close()
            init_db(old_db)
            local_id = create_event(old_db, "Local", "Host", "2026-08-25", "SCM", 6)
            self.assertEqual(get_event(old_db, local_id)["meet_type"], "Local")


if __name__ == "__main__":
    unittest.main()
