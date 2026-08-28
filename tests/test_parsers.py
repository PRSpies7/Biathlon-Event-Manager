from pathlib import Path
import io
import json

from openpyxl import Workbook, load_workbook

from exporters.filenames import event_filename, event_results_filename
from exporters.master_excel import build_master_import_xlsx
from exporters.swim_timekeeper import build_swim_timekeeper_xlsx
from exporters.timedrops_json import _timedrops_swimmer_name, generate_timedrops_json
from parsers.master_entries import parse_master_entries
from parsers.run_results_excel import parse_run_results_excel
from parsers.swim_results import parse_swim_results
from exporters.results_xml import build_results_xml

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"


def test_master_parser():
    x = parse_master_entries(DATA / "Master Entries Excel.xlsx")
    assert len(x["athletes"]) == 136
    assert x["running_heats"] == list(range(1, 14))
    assert x["swimming_heats"] == list(range(1, 19))
    assert x["athletes"][0]["athlete_number"] == "3235"


def test_run_parser():
    x = parse_run_results_excel(DATA / "Run Results LG4.xlsx")
    assert x["imported_heat_numbers"] == [1,2,3,4,5,6,7,8,9,10,11,13]
    heat1 = next(b for b in x["blocks"] if b["heat_numbers"] == (1,))
    assert heat1["results"][0] == {"position": 1, "run_time": "01:18.14", "row": 5}


def test_swim_parser_and_revision():
    x = parse_swim_results(DATA / "Swim TXT file.txt")
    assert x["athlete_results"]["7276"]["swim_time"] == "00:41.68"
    assert x["athlete_results"]["6702"]["swim_time"] == "03:55.05"
    assert x["revisions"]


def test_swim_parser_completely_ignores_no_start_entries():
    payload = io.BytesIO(
        b"Event #1 Heat 1 Race 1\n"
        b"1 1 Started Athlete (1001) 01:05.20 OK\n"
        b"2 0 NS 0.00\n"
        b"3 0 No Start Athlete (1002) NS\n"
    )
    parsed = parse_swim_results(payload)
    assert parsed["source_result_count"] == 1
    assert [r["athlete_number"] for r in parsed["result_records"]] == ["1001"]
    assert "1002" not in parsed["athlete_results"]


def test_xml_shape():
    payload = [{"athlete_number":"1011","athlete_name":"John Doe","run_time":"04:12.50","swim_time":"01:05.20"}]
    xml = build_results_xml(payload).decode("utf-8")
    assert xml.startswith('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>')
    assert '<results xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">' in xml
    assert '<athleteNo>1011</athleteNo>' in xml
    assert '<athleteSurname>John Doe</athleteSurname>' in xml


def test_run_parser_recognizes_combined_heat_labels():
    wb = Workbook()
    ws = wb.active
    ws.title = "Run times"
    ws.append(["Heat 11 and 12"])
    ws.append(["Lap", "Time"])
    ws.append([1, "01:20.10"])
    ws.append([2, "01:21.20"])
    payload = io.BytesIO()
    wb.save(payload)
    payload.seek(0)

    parsed = parse_run_results_excel(payload)
    assert parsed["imported_heat_numbers"] == [11, 12]
    assert parsed["blocks"][0]["heat_numbers"] == (11, 12)
    assert len(parsed["blocks"][0]["results"]) == 2


def test_run_parser_uses_first_worksheet_regardless_of_name():
    wb = Workbook()
    ws = wb.active
    ws.title = "Race Results"
    ws.append(["Heat 1"])
    ws.append(["Lap", "Time"])
    ws.append([1, "01:20.10"])
    wb.create_sheet("Run times")
    payload = io.BytesIO()
    wb.save(payload)
    payload.seek(0)

    parsed = parse_run_results_excel(payload)
    assert parsed["blocks"][0]["label"] == "Heat 1"


def test_uploaded_athlete_name_removes_afl_marker():
    wb = Workbook()
    ws = wb.active
    ws.append(["Running Heats"])
    ws.append(["Heat 1 - Test"])
    ws.append([1001, "Marieke Bouwer (AFL)", "U/19 WOMEN", 1])
    ws.append(["Swimming Heats"])
    ws.append(["Heat 1 - Test"])
    ws.append([1001, "Marieke Bouwer (AFL)", "U/19 WOMEN", 1])
    payload = io.BytesIO()
    wb.save(payload)
    payload.seek(0)

    parsed = parse_master_entries(payload)
    assert parsed["athletes"][0]["athlete_name"] == "Marieke Bouwer"


def test_timedrops_swimmer_name_keeps_id_within_30_characters():
    reference = json.loads((DATA / "TimeDrops JSON example.json").read_text(encoding="utf-8"))
    athlete = {
        "athlete_number": "7409", "athlete_name": "Christopher Alexander Smith", "group_name": "U/19 WOMEN",
        "swimming_heat": 1, "swimming_lane": 1,
    }
    generated = generate_timedrops_json(
        reference, [athlete], meet_name="Test", host_team="GN", start_date="2026-08-28",
        course="LCM", pool_lanes=8,
    )
    swimmer_name = generated["meetSwimmers"][0]["swimmerName"]
    assert swimmer_name == "Christopher Alexander S (7409)"
    assert len(swimmer_name) == 30
    assert swimmer_name.endswith(" (7409)")
    assert generated["meetSwimmers"][0]["swimmerId"] == "7409"
    assert _timedrops_swimmer_name("Elke Vorster", "7409") == "Elke Vorster (7409)"
    assert _timedrops_swimmer_name("Marieke Bouwer (AFL)", "7409") == "Marieke Bouwer (7409)"
    assert _timedrops_swimmer_name("Jürgen Müller", "7409") == "Jurgen Muller (7409)"


def test_swim_timekeeper_heat_rows_are_30_points_even_when_empty():
    data = build_swim_timekeeper_xlsx(
        [{"sort_order": 1, "athlete_number": "1001", "athlete_name": "Athlete", "group_name": "U/19 WOMEN", "swimming_heat": 1, "swimming_lane": 1},
         {"sort_order": 2, "athlete_number": "1002", "athlete_name": "Other", "group_name": "U/19 WOMEN", "swimming_heat": 3, "swimming_lane": 1}],
        "Test", 2,
    )
    wb = load_workbook(io.BytesIO(data.getvalue()))
    # Lane 2 has no swimmers, but it still has one row for each planned heat.
    assert wb["Lane2"].row_dimensions[8].height == 30
    assert wb["Lane2"].row_dimensions[9].height == 30


def test_master_excel_duplicates_athlete_name_into_athlete_surname():
    data = build_master_import_xlsx([
        {"athlete_number": "1011", "athlete_name": "Elke Vorster", "run_time": "04:12.50", "swim_time": "01:05.20"}
    ])
    ws = load_workbook(io.BytesIO(data.getvalue())).active
    assert [cell.value for cell in ws[1]] == ["Athlete nr", "Athlete Name", "athleteSurname", "runtime", "swimtime"]
    assert ws["B2"].value == ws["C2"].value == "Elke Vorster"


def test_event_results_filename_uses_a_sanitised_event_name():
    assert event_results_filename("Gauteng North Championships 2026", "xlsx") == "Gauteng North Championships 2026 Master Results.xlsx"
    assert event_results_filename('GN: Finals / 2026', ".xml") == "GN_ Finals _ 2026 Master Results.xml"
    assert event_filename('GN: Finals / 2026', "Athlete List", "xlsx") == "GN_ Finals _ 2026 Athlete List.xlsx"
