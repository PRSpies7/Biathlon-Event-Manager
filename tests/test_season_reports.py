from copy import deepcopy
from dataclasses import replace
from datetime import date
from io import BytesIO

import openpyxl
import pytest

from database.season_db import init_season_db
from database.season_repository import replace_record_benchmarks, reset_season_database, season_snapshot, store_event
from exporters.season_results import build_report_xlsx, build_season_results_xlsx, build_qualified_athletes_xlsx
from parsers.season_records import RECORD_HEADERS, parse_record_benchmarks
from parsers.season_results.models import INTERPROVINCIAL, LEAGUE
from services.season_reports import (
    RECORD_REPORT_HEADERS, completed_result, filtered_snapshot,
    qualification_rows, record_comparison_rows, school_snapshot, age_group_sort_key,
    athlete_result_rows, qualified_report_rows, report_sort_key, top_athlete_report,
)
from test_season_results import sample_event


@pytest.fixture
def db(tmp_path):
    path = tmp_path / "reports.sqlite"
    init_season_db(path)
    return path


def add_events(db, category="U/11 GIRLS", leagues=3, interprovincials=1):
    base = sample_event()
    result = replace(base.results[0], category=category, affiliated=False)
    for i, kind in enumerate([LEAGUE] * leagues + [INTERPROVINCIAL] * interprovincials):
        store_event(db, replace(base, name=f"Event {i+1}", event_date=date(2026, 8, i+1),
                                competition_type=kind, results=[result], awards=[]))
    return season_snapshot(db)


@pytest.mark.parametrize("category,leagues,ip,qualified,required", [
    ("U/11 GIRLS", 3, 1, "Yes", 3), ("U/19 BOYS", 2, 1, "No", 3),
    ("U/19 BOYS", 3, 0, "No", 3), ("U/11 GIRLS", 4, 2, "Yes", 3),
    ("JNR WOMEN", 2, 1, "Yes", 2), ("JUNIOR MEN", 2, 1, "Yes", 2),
    ("JUNIORS MEN", 1, 1, "No", 2), ("SENIORS MEN", 2, 1, "Yes", 2),
    ("SENIOR WOMEN", 2, 0, "No", 2), ("MASTERS 40+ MEN", 2, 1, "Yes", 2),
    ("MASTER 80+ MEN", 2, 1, "Yes", 2), ("SPECIAL NEEDS MALE", 2, 1, "No", 3),
])
def test_qualification_rules(db, category, leagues, ip, qualified, required):
    row = qualification_rows(add_events(db, category, leagues, ip))[0]
    assert row["Qualified"] == qualified
    assert row["Completed league events"] == leagues
    assert row["Completed interprovincial events"] == ip
    assert row["Required completed events"] == required + 1
    assert row["Affiliated"] == "Not Affiliated"  # No extra affiliation eligibility rule.
    assert "Event 1 (2026-08-01)" in row["Counted events"]


@pytest.mark.parametrize("field,value", [("status", "DNF"), ("status", "DNS"), ("status", "DSQ"),
    ("run_time", "00:00.000"), ("swim_time", "00:00.00"), ("run_time", ""), ("swim_time", "DNF")])
def test_incomplete_events_do_not_count(db, field, value):
    snapshot = add_events(db)
    snapshot["results"][0][field] = value
    assert not completed_result(snapshot["results"][0])
    row = qualification_rows(snapshot)[0]
    assert row["Completed league events"] == 2 and row["Qualified"] == "No"


def test_unique_event_count_and_chronological_category(db):
    snapshot = add_events(db, category="SENIORS MEN", leagues=2)
    snapshot["athletes"][0]["category"] = "U/19 BOYS"
    snapshot["results"].append(deepcopy(snapshot["results"][0]))
    row = qualification_rows(snapshot)[0]
    assert row["Completed league events"] == 2
    assert row["Category"] == "SENIORS MEN" and row["Qualified"] == "Yes"


def test_school_and_affiliation_filters_preserve_known_school_without_mutation(db):
    event = sample_event()
    store_event(db, event)
    later = replace(event, name="IP", competition_type=INTERPROVINCIAL,
                    results=[replace(event.results[0], school="", affiliated=False), event.results[1]])
    store_event(db, later)
    snapshot = season_snapshot(db)
    before = deepcopy(snapshot)
    schools = school_snapshot(snapshot)
    assert len(schools["athletes"]) == 1 and len(schools["results"]) == 2
    assert all(r["school"] == "Example School" for r in schools["results"])
    assert all(e["result_count"] == 1 for e in schools["events"])
    affiliated = filtered_snapshot(snapshot, {a["athlete_number"] for a in snapshot["athletes"] if a["affiliated"]})
    assert len(affiliated["athletes"]) == 1 and len(affiliated["awards"]) == 6
    assert snapshot == before
    for report in (schools, affiliated):
        workbook = openpyxl.load_workbook(build_season_results_xlsx(report))
        assert workbook["Athletes"].max_row == 2 and workbook["Event Results"].max_row == 3
        athlete_rows = list(workbook["Athletes"].values)
        assert athlete_rows[1][athlete_rows[0].index("Affiliated")] == "Affiliated"
        result_rows = list(workbook["Event Results"].values)
        exported = [dict(zip(result_rows[0], row)) for row in result_rows[1:]]
        assert {r["Event name"] for r in exported} == {"GN League 1", "IP"}
        assert all(r["Event date"].date() == date(2026, 8, 25) for r in exported)
        assert all(r["Run time"] == "01:20.10" and r["Swim time"] == "00:40.50"
                   and r["Total points"] == 2060 for r in exported)
        workbook.close()


def record_reference(category="U/11 GIRLS", points="2050.00"):
    return {"category": category, "athlete_number": "00999", "athlete_name": "Record Holder", "total_points": points}


def test_record_template_parse_persistence_replace_reset_and_full_export(db):
    template = build_report_xlsx("Current Records", RECORD_HEADERS, [
        {"Category": "U/11 GIRLS", "Athlete number": "00999", "Athlete name": "Record Holder", "Total points": 2050},
        {"Category": "U/13 BOYS"},
    ])
    references = parse_record_benchmarks(template)
    assert len(references) == 1 and references[0]["athlete_number"] == "00999"
    replace_record_benchmarks(db, references, "records.xlsx")
    init_season_db(db)
    snapshot = season_snapshot(db)
    assert len(snapshot["records"]) == 1
    assert not snapshot["athletes"]  # Record holders need not be season participants.
    workbook = openpyxl.load_workbook(build_season_results_xlsx(snapshot))
    assert workbook["Current Records"]["B2"].value == "00999"
    assert workbook["Current Records"]["D2"].value == 2050
    workbook.close()
    replace_record_benchmarks(db, [record_reference("JNR MEN")], "new.xlsx")
    assert [r["category"] for r in season_snapshot(db)["records"]] == ["JNR MEN"]
    reset_season_database(db)
    assert not season_snapshot(db)["records"]


@pytest.mark.parametrize("points,comparison", [("2059.99", "Exceeded"), ("2060.000", "Equalled"), ("2060.01", None)])
def test_category_records_exact_published_points(db, points, comparison):
    store_event(db, sample_event())
    replace_record_benchmarks(db, [record_reference("U11 GIRLS", points)], "records.xlsx")
    rows = record_comparison_rows(season_snapshot(db))
    assert len(rows) == (1 if comparison else 0)
    if comparison:
        assert rows[0]["Record comparison"] == comparison
        assert rows[0]["Reference athlete number"] == "00999"
        assert rows[0]["Athlete number"] == "00101"
        assert rows[0]["Total points"] == 2060
        workbook = openpyxl.load_workbook(build_report_xlsx("Records", RECORD_REPORT_HEADERS, rows))
        assert workbook.active["F2"].value == 2060
        workbook.close()


def test_records_compare_category_and_published_points_only(db):
    store_event(db, sample_event())
    replace_record_benchmarks(db, [record_reference("U13 BOYS", "100")], "records.xlsx")
    assert not record_comparison_rows(season_snapshot(db))
    replace_record_benchmarks(db, [record_reference()], "records.xlsx")
    snapshot = season_snapshot(db)
    snapshot["results"][0]["status"] = "DNF"
    assert record_comparison_rows(snapshot)


@pytest.mark.parametrize("references", [[record_reference(points="DNF")], [record_reference(points="NaN")],
    [record_reference(points="-1")], [record_reference(), record_reference("U11 GIRLS")],
    [dict(record_reference(), category="")]])
def test_invalid_records_leave_saved_references_intact(db, references):
    replace_record_benchmarks(db, [record_reference()], "records.xlsx")
    before = season_snapshot(db)
    with pytest.raises(ValueError):
        replace_record_benchmarks(db, references, "invalid.xlsx")
    assert season_snapshot(db) == before


def test_qualified_workbook_matches_top_athletes_with_status_and_shortfall(db):
    snapshot = add_events(db)
    base_athlete = snapshot["athletes"][0]
    base_result = snapshot["results"][0]
    snapshot["athletes"].append({**base_athlete, "id": 2, "athlete_number": "00202", "athlete_name": "Pending Athlete"})
    snapshot["results"].append({**base_result, "athlete_id": 2, "athlete_number": "00202", "athlete_name": "Pending Athlete"})
    top_headers, top_rows = top_athlete_report(snapshot)
    rows = qualified_report_rows(snapshot)
    assert [{h: row[h] for h in top_headers[:10]} for row in rows] == [{h: row[h] for h in top_headers[:10]} for row in top_rows]
    assert rows[0]["Qualified"] == "Yes" and rows[0]["Missing to qualify"] == "None"
    assert rows[1]["Qualified"] == "No"
    assert rows[1]["Missing to qualify"] == "3 more completed events; including at least 1 interprovincial or Gauteng North Championship"
    workbook = openpyxl.load_workbook(build_qualified_athletes_xlsx(snapshot, rows))
    assert workbook.sheetnames == ["Qualified Athletes", "Insights"]
    exported = list(workbook.active.values)
    expected_headers = list(top_headers[:10])
    expected_headers.insert(expected_headers.index("Athlete name") + 1, "Qualified")
    expected_headers.append("Missing to qualify")
    assert list(exported[0]) == expected_headers
    assert workbook.active.max_row == 3 and workbook.active.freeze_panes == "E2"
    assert exported[1][3] == "Yes" and exported[1][-1] == "None"
    assert exported[2][3] == "No" and exported[2][-1] == "3 more completed events; including at least 1 interprovincial or Gauteng North Championship"
    workbook.close()


@pytest.mark.parametrize("category,leagues,ip,missing", [
    ("U/11 GIRLS", 2, 1, "1 more completed event"),
    ("U/11 GIRLS", 3, 0, "1 more completed event; including at least 1 interprovincial or Gauteng North Championship"),
    ("JNR WOMEN", 1, 1, "1 more completed event"),
    ("MASTERS 60+ MEN", 2, 1, "None"),
])
def test_qualification_report_shortfalls_use_completed_event_requirements(db, category, leagues, ip, missing):
    rows = qualified_report_rows(add_events(db, category, leagues, ip))
    assert len(rows) == 1
    assert rows[0]["Missing to qualify"] == missing
    assert rows[0]["Qualified"] == ("Yes" if missing == "None" else "No")


@pytest.mark.parametrize("category,leagues,ip,qualified", [
    ("U/11 GIRLS", 2, 2, "Yes"), ("U/19 BOYS", 1, 3, "Yes"),
    ("U/11 GIRLS", 4, 0, "No"), ("JNR WOMEN", 1, 2, "Yes"),
    ("SENIORS MEN", 0, 3, "Yes"), ("MASTERS 60+ WOMEN", 3, 0, "No"),
])
def test_individual_qualification_allows_mixed_meet_types(db, category, leagues, ip, qualified):
    assert qualified_report_rows(add_events(db, category, leagues, ip))[0]["Qualified"] == qualified


def test_individual_championship_only_counts_once(db):
    snapshot = add_events(db, leagues=2, interprovincials=1)
    snapshot["results"][-1]["event_name"] = "GN Championship"
    assert qualification_rows(snapshot)[0]["Completed distinct qualifying events"] == 3
    assert qualified_report_rows(snapshot)[0]["Qualified"] == "No"
    for result in snapshot["results"]:
        result["category"] = "JNR WOMEN"
    assert qualified_report_rows(snapshot)[0]["Qualified"] == "Yes"


def test_natural_age_order_then_descending_numeric_score():
    groups = ["MASTERS 80+ MEN", "SPECIAL NEEDS MALE", "U/19 BOYS", "SENIORS MEN", "U/08 GIRLS",
              "JNR WOMEN", "MASTERS 40+ WOMEN", "U/11 GIRLS", "U/09 BOYS"]
    assert sorted(groups, key=age_group_sort_key) == ["U/08 GIRLS", "U/09 BOYS", "U/11 GIRLS", "U/19 BOYS",
        "JNR WOMEN", "SENIORS MEN", "MASTERS 40+ WOMEN", "MASTERS 80+ MEN", "SPECIAL NEEDS MALE"]
    rows = [{"Age group": "U/11 GIRLS", "Total points": p} for p in ["999", "DNF", "2100", "0", "1000"]]
    assert [r["Total points"] for r in sorted(rows, key=report_sort_key)] == ["2100", "1000", "999", "0", "DNF"]


def test_summary_times_and_total_come_from_same_best_event(db):
    snapshot = add_events(db)
    snapshot["results"][1].update(total_points="2200", run_time="01:11.00", swim_time="00:32.00")
    summary = athlete_result_rows(snapshot)[0]
    assert summary["Total points"] == 2200
    assert summary["Run time"] == "01:11.00" and summary["Swim time"] == "00:32.00"
    assert summary["Result event"] == "Event 2 (2026-08-02)"
    assert summary["School"] == "Example School"
    snapshot["results"][-1]["category"] = "U/13 GIRLS"
    summary = athlete_result_rows(snapshot)[0]
    assert summary["Age group"] == "U/13 GIRLS" and summary["Total points"] == 2060
    assert "Event 4" in summary["Result event"]


def test_affiliated_export_has_no_points_anywhere(db):
    snapshot = add_events(db)
    workbook = openpyxl.load_workbook(build_season_results_xlsx(snapshot, include_points=False))
    for sheet in workbook:
        assert all("points" not in str(cell.value).casefold() for cell in sheet[1])
    headers = [c.value for c in workbook["Athletes"][1]]
    assert {"Age group", "School", "Run time", "Swim time"}.issubset(headers)
    workbook.close()


def test_full_and_school_result_sheets_are_sorted_by_age_and_points(db):
    snapshot = add_events(db)
    snapshot["results"][0].update(category="MASTERS 40+ MEN", total_points="3000")
    snapshot["results"][1].update(category="U/08 GIRLS", total_points="1000")
    snapshot["results"][2].update(category="U/08 GIRLS", total_points="2100")
    snapshot["results"][3].update(category="U/08 GIRLS", total_points="DNF")
    for filtered in (snapshot, school_snapshot(snapshot)):
        workbook = openpyxl.load_workbook(build_season_results_xlsx(filtered))
        rows = list(workbook["Event Results"].values)
        headers = rows[0]
        points_index = headers.index("Total points")
        assert [r[points_index] for r in rows[1:]] == [2100, 1000, "DNF", 3000]
        assert {"Athlete number", "Athlete name", "Age group", "School", "Run time", "Swim time"}.issubset(headers)
        workbook.close()


def test_top_athletes_ranks_all_by_best_score_with_ties_and_event_totals(db):
    snapshot = add_events(db)
    base_athlete = snapshot["athletes"][0]
    base_result = snapshot["results"][0]
    for aid, name, score in [(2, "Second Athlete", "2100"), (3, "Third Athlete", "2100"), (4, "No Finish", "DNF")]:
        snapshot["athletes"].append({**base_athlete, "id": aid, "athlete_number": str(aid), "athlete_name": name,
                                     "affiliated": aid == 2})
        snapshot["results"].append({**base_result, "athlete_id": aid, "athlete_number": str(aid), "athlete_name": name, "total_points": score})
    headers, rows = top_athlete_report(snapshot)
    assert [r["Athlete name"] for r in rows] == ["Second Athlete", "Third Athlete", "Alex Example", "No Finish"]
    assert [r["Rank"] for r in rows] == [1, 1, 3, ""]
    assert rows[0]["Highest total points"] == 2100
    assert rows[0]["Event 1 (2026-08-01)"] == 2100 and rows[0]["Event 2 (2026-08-02)"] == ""
    assert rows[-1]["Event 1 (2026-08-01)"] == "DNF"
    assert rows[2]["Events attended"] == 4
    assert not any("time" in h.lower() for h in headers)
    workbook = openpyxl.load_workbook(build_report_xlsx("Top Athletes", headers, rows))
    assert workbook.active.max_row == 5 and workbook.active.freeze_panes == "E2"
    assert workbook.active["K2"].value == 2100
    exported = list(workbook.active.values)
    affiliation_column = exported[0].index("Affiliated")
    assert exported[1][affiliation_column] == "Affiliated"
    assert exported[2][affiliation_column] == "Not Affiliated"
    workbook.close()


def test_record_report_includes_school_times_age_and_is_score_sorted(db):
    snapshot = add_events(db)
    replace_record_benchmarks(db, [record_reference(points="2000")], "records.xlsx")
    snapshot["records"] = season_snapshot(db)["records"]
    snapshot["results"][1]["total_points"] = "2200"
    rows = record_comparison_rows(snapshot)
    assert rows[0]["Total points"] == 2200
    assert rows[0]["Age group"] == "U/11 GIRLS"
    assert rows[0]["School"] == "Example School"
    assert rows[0]["Run time"] == "01:20.10" and rows[0]["Swim time"] == "00:40.50"
    workbook = openpyxl.load_workbook(build_report_xlsx("Records", RECORD_REPORT_HEADERS, rows))
    exported = list(workbook.active.values)
    record = dict(zip(exported[0], exported[1]))
    assert record["Event"] == "Event 2"
    assert record["Event date"].date() == date(2026, 8, 2)
    workbook.close()
