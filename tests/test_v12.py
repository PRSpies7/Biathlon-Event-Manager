from pathlib import Path

from openpyxl import load_workbook

from database.db import init_db
from database.repository import (
    add_athlete_to_run_heat,
    assign_athlete_to_run_heat,
    create_event,
    get_athletes,
    get_run_groups,
    get_running_heats,
    replace_athletes,
    remove_athlete_from_run_heat,
    restore_run_position_mappings,
    save_run_positions,
)
from exporters.athlete_event_mapping import build_printable_athlete_mapping_xlsx
from exporters.run_positions import build_run_positions_xlsx
from parsers.run_positions_excel import parse_run_positions_excel
from parsers.master_entries import parse_master_entries
from parsers.run_results_excel import parse_run_results_excel
from parsers.swim_results import parse_swim_results
from services.results_service import process_run_results, process_swim_results
from validation.validators import normalize_time

BASE = Path(__file__).resolve().parents[1]


def test_v12_time_normalization_variants():
    assert normalize_time("01.01.20") == "01:01.20"
    assert normalize_time("01,01,20") == "01:01.20"
    assert normalize_time("01;01;20") == "01:01.20"
    assert normalize_time("1:05.20") == "01:05.20"


def test_v12_run_processing_is_heat_aware_and_persists(tmp_path):
    parsed = parse_master_entries(BASE / "data" / "Master Entries Excel.xlsx")
    db = tmp_path / "event.sqlite"
    init_db(db)
    event_id = create_event(str(db), "Test", "GN", "2026-08-14", "LCM", 8)
    replace_athletes(str(db), event_id, parsed["athletes"])

    athletes = get_athletes(str(db), event_id)
    # Create the Phase 2 mapping required for every heat in the supplied run file.
    for heat in parsed["running_heats"]:
        group = [a for a in athletes if a.get("running_heat") == heat]
        save_run_positions(str(db), event_id, [heat], {a["athlete_number"]: i + 1 for i, a in enumerate(group)})

    run_data = parse_run_results_excel(BASE / "data" / "Run Results LG4.xlsx")
    result = process_run_results(str(db), event_id, run_data)

    assert result["saved"] is True
    assert result["matched_count"] > 0
    assert any("missing heat section(s): 12" in w for w in result["warnings"])

    persisted = get_athletes(str(db), event_id)
    assert sum(bool(a.get("run_time")) for a in persisted) == result["matched_count"]


def test_v12_swim_processing_is_independent(tmp_path):
    parsed = parse_master_entries(BASE / "data" / "Master Entries Excel.xlsx")
    db = tmp_path / "event.sqlite"
    init_db(db)
    event_id = create_event(str(db), "Test", "GN", "2026-08-14", "LCM", 8)
    replace_athletes(str(db), event_id, parsed["athletes"])

    swim_data = parse_swim_results(BASE / "data" / "Swim TXT file.txt")
    result = process_swim_results(str(db), event_id, swim_data)

    assert result["saved"] is True
    assert result["matched_count"] > 0
    persisted = get_athletes(str(db), event_id)
    assert sum(bool(a.get("swim_time")) for a in persisted) == result["matched_count"]


def test_v12_combined_run_heat_group_persists_all_results(tmp_path):
    db = tmp_path / "combined.sqlite"
    init_db(db)
    event_id = create_event(str(db), "Test", "GN", "2026-08-14", "LCM", 8)
    replace_athletes(str(db), event_id, [
        {"sort_order": 1, "athlete_number": "1101", "athlete_name": "A", "group_name": "U/19 MEN", "running_heat": 11, "running_lane": 1, "swimming_heat": 1, "swimming_lane": 1},
        {"sort_order": 2, "athlete_number": "1201", "athlete_name": "B", "group_name": "U/19 MEN", "running_heat": 12, "running_lane": 2, "swimming_heat": 1, "swimming_lane": 2},
    ])
    save_run_positions(str(db), event_id, [11, 12], {"1101": 1, "1201": 2})

    run_data = {
        "blocks": [{
            "label": "Heat 11 and 12",
            "heat_numbers": (11, 12),
            "results": [
                {"position": 1, "run_time": "01:20.10"},
                {"position": 2, "run_time": "01:21.20"},
            ],
        }],
        "imported_heat_numbers": [11, 12],
    }
    result = process_run_results(str(db), event_id, run_data)
    assert result["matched_count"] == 2
    persisted = {a["athlete_number"]: a for a in get_athletes(str(db), event_id)}
    assert persisted["1101"]["run_time"] == "01:20.10"
    assert persisted["1201"]["run_time"] == "01:21.20"
    assert any("Combined run heat label(s) detected" in w for w in result["warnings"])


def test_existing_athlete_can_move_to_another_run_heat_without_duplication(tmp_path):
    db = tmp_path / "event.sqlite"
    init_db(db)
    event_id = create_event(str(db), "Test", "GN", "2026-08-14", "LCM", 8)
    replace_athletes(str(db), event_id, [
        {"sort_order": 1, "athlete_number": "1101", "athlete_name": "A", "group_name": "U/19 MEN", "running_heat": 1, "running_lane": 1, "swimming_heat": 1, "swimming_lane": 1},
        {"sort_order": 2, "athlete_number": "1102", "athlete_name": "B", "group_name": "U/19 MEN", "running_heat": 2, "running_lane": 1, "swimming_heat": 1, "swimming_lane": 2},
    ])
    assert assign_athlete_to_run_heat(str(db), event_id, "1101", 2) is True
    assert assign_athlete_to_run_heat(str(db), event_id, "1101", 2) is False
    athletes = get_athletes(str(db), event_id)
    assert len(athletes) == 2
    moved = next(a for a in athletes if a["athlete_number"] == "1101")
    assert moved["running_heat"] == 2
    assert moved["run_position"] is None


def test_new_numbered_athlete_is_added_to_selected_run_heat(tmp_path):
    db = tmp_path / "event.sqlite"
    init_db(db)
    event_id = create_event(str(db), "Test", "GN", "2026-08-14", "LCM", 8)
    replace_athletes(str(db), event_id, [
        {"sort_order": 1, "athlete_number": "1101", "athlete_name": "A", "group_name": "U/19 MEN", "running_heat": 1, "running_lane": 1, "swimming_heat": 1, "swimming_lane": 1},
    ])
    add_athlete_to_run_heat(str(db), event_id, "9999", "New Athlete", 1)
    athletes = {a["athlete_number"]: a for a in get_athletes(str(db), event_id)}
    assert athletes["9999"]["athlete_name"] == "New Athlete"
    assert athletes["9999"]["running_heat"] == 1
    assert athletes["9999"]["swimming_heat"] is None
    try:
        add_athlete_to_run_heat(str(db), event_id, "9999", "Duplicate", 1)
    except ValueError as exc:
        assert "already exists" in str(exc)
    else:
        raise AssertionError("A duplicate athlete number was accepted")
    assert len(get_athletes(str(db), event_id)) == 2


def test_remove_athlete_clears_only_run_assignment_and_preserves_empty_heat(tmp_path):
    db = tmp_path / "event.sqlite"
    init_db(db)
    event_id = create_event(str(db), "Test", "GN", "2026-08-14", "SCM", 8)
    replace_athletes(str(db), event_id, [
        {"sort_order": 1, "athlete_number": "0011", "athlete_name": "Athlete A", "group_name": "U/19 MEN", "running_heat": 7, "running_lane": 2, "swimming_heat": 3, "swimming_lane": 4},
    ])
    save_run_positions(str(db), event_id, [7], {"0011": 1})

    assert remove_athlete_from_run_heat(str(db), event_id, "0011", [7]) is True
    athlete = get_athletes(str(db), event_id)[0]
    assert athlete["athlete_number"] == "0011"
    assert athlete["athlete_name"] == "Athlete A"
    assert athlete["running_heat"] is None
    assert athlete["running_lane"] is None
    assert athlete["run_group_key"] is None
    assert athlete["run_position"] is None
    assert athlete["swimming_heat"] == 3
    assert athlete["swimming_lane"] == 4
    assert get_running_heats(str(db), event_id) == [7]


def test_printable_athlete_mapping_is_sorted_and_configured_for_portrait_printing():
    workbook = build_printable_athlete_mapping_xlsx([
        {"athlete_number": "20", "athlete_name": "Zulu Athlete", "group_name": "Senior", "running_heat": 2, "swimming_heat": 4, "swimming_lane": 5},
        {"athlete_number": "0011", "athlete_name": "Alpha Athlete", "group_name": "Junior", "running_heat": 1, "swimming_heat": 3, "swimming_lane": 2},
    ])
    ws = load_workbook(workbook).active

    assert [cell.value for cell in ws[1]] == [
        "Athlete Number", "Athlete Name", "Age Group", "Run Heat", "Swim Heat", "Swim Lane"
    ]
    assert [ws.cell(2, column).value for column in range(1, 7)] == [
        "0011", "Alpha Athlete", "Junior", 1, 3, 2
    ]
    assert ws["A2"].number_format == "@"
    assert ws.row_dimensions[2].height == 19
    assert ws.page_setup.orientation == "portrait"
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 0
    assert ws.print_title_rows == "$1:$1"


def test_run_position_export_import_restores_group_and_positions(tmp_path):
    original_db = tmp_path / "original.sqlite"
    restored_db = tmp_path / "restored.sqlite"
    for db in (original_db, restored_db):
        init_db(db)
    original_event = create_event(str(original_db), "Test", "GN", "2026-08-14", "LCM", 8)
    restored_event = create_event(str(restored_db), "Test", "GN", "2026-08-14", "LCM", 8)
    seed = [
        {"sort_order": 1, "athlete_number": "1101", "athlete_name": "A", "group_name": "U/19 MEN", "running_heat": 1, "running_lane": 1, "swimming_heat": 1, "swimming_lane": 1},
        {"sort_order": 2, "athlete_number": "1102", "athlete_name": "B", "group_name": "U/19 MEN", "running_heat": 2, "running_lane": 1, "swimming_heat": 1, "swimming_lane": 2},
    ]
    replace_athletes(str(original_db), original_event, seed)
    restored_seed = [dict(seed[0], running_heat=2), dict(seed[1], running_heat=1)]
    replace_athletes(str(restored_db), restored_event, restored_seed)
    save_run_positions(str(original_db), original_event, [1], {"1101": 1})
    save_run_positions(str(original_db), original_event, [2], {"1102": 1})

    exported = build_run_positions_xlsx(
        get_athletes(str(original_db), original_event),
        get_run_groups(str(original_db), original_event),
    )
    restored_count = restore_run_position_mappings(
        str(restored_db), restored_event, parse_run_positions_excel(exported)
    )
    athletes = {a["athlete_number"]: a for a in get_athletes(str(restored_db), restored_event)}
    assert restored_count == 2
    assert athletes["1101"]["running_heat"] == 1
    assert athletes["1101"]["run_group_key"] == "1"
    assert athletes["1101"]["run_position"] == 1
    assert athletes["1102"]["running_heat"] == 2
    assert athletes["1102"]["run_group_key"] == "2"
    assert athletes["1102"]["run_position"] == 1


def test_swim_results_match_by_id_then_safe_truncated_name_fallback(tmp_path):
    db = tmp_path / "event.sqlite"
    init_db(db)
    event_id = create_event(str(db), "Test", "GN", "2026-08-14", "LCM", 8)
    replace_athletes(str(db), event_id, [
        {"sort_order": 1, "athlete_number": "1101", "athlete_name": "Normal Athlete", "group_name": "U/19 MEN", "running_heat": 1, "running_lane": 1, "swimming_heat": 1, "swimming_lane": 1},
        {"sort_order": 2, "athlete_number": "1102", "athlete_name": "Marieke Bouwer (AFL)", "group_name": "U/19 WOMEN", "running_heat": 1, "running_lane": 2, "swimming_heat": 1, "swimming_lane": 2},
        {"sort_order": 3, "athlete_number": "1103", "athlete_name": "A Very Long Athlete Name That Continues", "group_name": "U/19 WOMEN", "running_heat": 1, "running_lane": 3, "swimming_heat": 1, "swimming_lane": 3},
        {"sort_order": 4, "athlete_number": "1104", "athlete_name": "Different Athlete", "group_name": "U/19 WOMEN", "running_heat": 1, "running_lane": 4, "swimming_heat": 1, "swimming_lane": 4},
        {"sort_order": 5, "athlete_number": "1105", "athlete_name": "Jürgen Müller", "group_name": "U/19 MEN", "running_heat": 1, "running_lane": 5, "swimming_heat": 1, "swimming_lane": 5},
        {"sort_order": 6, "athlete_number": "1106", "athlete_name": "Shared Long Athlete Name One", "group_name": "U/19 MEN", "running_heat": 1, "running_lane": 6, "swimming_heat": 1, "swimming_lane": 6},
        {"sort_order": 7, "athlete_number": "1107", "athlete_name": "Shared Long Athlete Name Two", "group_name": "U/19 MEN", "running_heat": 1, "running_lane": 7, "swimming_heat": 1, "swimming_lane": 7},
    ])
    result = process_swim_results(str(db), event_id, {
        "result_records": [
            {"athlete_number": "1101", "athlete_name": "Wrong Name", "swim_time": "01:01.00"},
            {"athlete_number": "", "athlete_name": "  marieke   bouwer ", "swim_time": "01:02.00"},
            {"athlete_number": "999", "athlete_name": "A Very Long Athlete Name", "swim_time": "01:03.00"},
            {"athlete_number": "110", "athlete_name": "Jurgen Muller", "swim_time": "01:03.50"},
            {"athlete_number": "", "athlete_name": "Shared Long Athlete Name", "swim_time": "01:03.60"},
            {"athlete_number": "", "athlete_name": "No Such Athlete", "swim_time": "01:04.00"},
        ],
    })
    persisted = {a["athlete_number"]: a for a in get_athletes(str(db), event_id)}
    assert result["matched_count"] == 4
    assert persisted["1101"]["swim_time"] == "01:01.00"
    assert persisted["1102"]["swim_time"] == "01:02.00"
    assert persisted["1103"]["swim_time"] == "01:03.00"
    assert persisted["1104"]["swim_time"] is None
    assert persisted["1105"]["swim_time"] == "01:03.50"
    assert persisted["1106"]["swim_time"] is None
    assert persisted["1107"]["swim_time"] is None
    assert any("No Such Athlete" in issue for issue in result["issues"])
    assert any("matches multiple athletes" in issue for issue in result["issues"])
