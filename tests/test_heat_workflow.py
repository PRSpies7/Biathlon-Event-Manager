"""Representative feature rules; shared fixtures also exercise final exports."""
from dataclasses import replace
from datetime import date
from pathlib import Path
import json
import sqlite3
import pytest

from database.db import init_db
from database.season_db import init_season_db
from database.season_repository import store_event, season_snapshot
from parsers.season_results.models import NormalizedEvent, Result, LEAGUE
from services.competition import infer_season
from services.seeding import select_seed, prepare_entries
from services.heats import generate_heats, centre_out


def test_inferred_season_and_three_season_distance_lookup(tmp_path):
    assert [infer_season(d) for d in ["2026-07-31","2026-08-01","2026-09-01","2027-01-01"]] == [2026,2027,2027,2027]
    db = tmp_path / "history.sqlite"
    init_season_db(db)
    def save(year,name,time,distance=400):
        store_event(db,NormalizedEvent(name,date(year,1,1),LEAGUE,"sample.xlsx",
            [Result("00101","Alex Example","U/11 GIRLS",run_time=time,run_distance=distance)],season_year=year))
    save(2024,"Too old","00:10.00")
    save(2025,"Fallback","01:19.00")
    save(2026,"Previous","01:18.00")
    save(2027,"Current","01:21.00")
    save(2027,"Current best","01:20.00")
    save(2027,"Other distance","02:40.00",800)
    aid=season_snapshot(db)["athletes"][0]["id"]
    assert select_seed(db,aid,"run",400,2027)[0] == 8000
    assert select_seed(db,aid,"run",400,2028)[0] == 8000
    assert select_seed(db,aid,"run",400,2029)[0] == 8000
    assert select_seed(db,aid,"run",400,2030)[0] is None
    assert select_seed(db,aid,"run",800,2027)[0] == 16000
    assert select_seed(db,aid,"swim",50,2027)[0] is None


def test_entry_matching_uses_number_and_highlights_name_difference(tmp_path):
    db=tmp_path/"history.sqlite"
    init_season_db(db)
    store_event(db,NormalizedEvent("League",date(2026,9,1),LEAGUE,"x.xlsx",
        [Result("00101","Alex Example","U/11 GIRLS",run_time="01:20.00")],season_year=2027))
    parsed={"athletes":[{"athlete_number":"101","athlete_name":"Different Person","group_name":"U/13 GIRLS",
        "running_heat":1,"swimming_heat":1,"sort_order":1}]}
    rows,ambiguous=prepare_entries(parsed,db,2027)
    assert len(ambiguous)==1 and ambiguous[0]["number_matched"]
    assert rows[0]["history_athlete_id"] is not None
    assert rows[0]["athlete_name"]=="Different Person"
    assert rows[0]["run_seed"] is None  # U13 distance differs from U11 history.
    parsed["athletes"][0]["athlete_name"]="Alex Example"
    rows,ambiguous=prepare_entries(parsed,db,2027)
    assert not ambiguous and rows[0]["history_athlete_id"] is not None
    assert rows[0]["run_distance"]==800 and rows[0]["run_seed"] is None


def test_programme_sequence_protected_opening_and_reseeded_moves():
    from services.heats import move_or_swap
    import pytest
    groups=["U/19 GIRLS","U/17 GIRLS","U/15 GIRLS","U/13 GIRLS","MASTERS 50+ WOMEN",
        "MASTERS 40+ WOMEN","SENIOR WOMEN","JNR GIRLS","U/11 GIRLS","U/09 GIRLS","U/08 GIRLS",
        "SPECIAL NEEDS FEMALE","MASTERS 80+ WOMEN","MASTERS 70+ WOMEN","MASTERS 60+ WOMEN"]
    field=[entries(1,group,i)[0] for i,group in enumerate(groups)]
    rows=generate_heats(field,settings(),optimise=False)
    run=[r["group_name"] for r in sorted(rows,key=lambda r:r["running_heat"])]
    swim=[r["group_name"] for r in sorted(rows,key=lambda r:r["swimming_heat"])]
    assert run==list(reversed(groups))
    assert swim[:5]==["U/08 GIRLS",*run[:4]]
    assert swim.index("U/13 GIRLS")<swim.index("JNR GIRLS")<swim.index("U/15 GIRLS")
    field=entries(6,"MASTERS 60+ WOMEN")+entries(2,"MASTERS 70+ WOMEN",10)+entries(2,"SPECIAL NEEDS FEMALE",20)+entries(3,"U/08 GIRLS",30)
    rows=generate_heats(field,settings())
    opening=[r for r in rows if r["running_heat"]==1]
    assert len(opening)==10  # Efficient League absorbs the compatible Special Needs remainder.
    assert {r["group_name"] for r in opening}=={"MASTERS 60+ WOMEN","MASTERS 70+ WOMEN","SPECIAL NEEDS FEMALE"}
    field=generate_heats(entries(16),settings(),optimise=False)
    before={r["athlete_number"]:dict(r) for r in field}
    selected=field[0]
    destination=next(r["running_heat"] for r in field if r["running_heat"]!=selected["running_heat"])
    moved=move_or_swap(field,settings(),"run",selected["athlete_number"],destination)
    members=sorted([r for r in moved if r["running_heat"]==destination],key=lambda r:r["run_seed"])
    assert [r["running_lane"] for r in members]==list(range(len(members),0,-1))
    assert all(r["swimming_lane"]==before[r["athlete_number"]]["swimming_lane"] for r in moved)
    swim=generate_heats(entries(7),settings(),optimise=False)
    swapped=move_or_swap(swim,settings(),"swim",swim[0]["athlete_number"],1,swim[-1]["athlete_number"])
    assert swapped[0]["swimming_heat"]==swim[-1]["swimming_heat"]
    full=generate_heats(entries(7),settings())
    for i,row in enumerate(full):
        row.update(swimming_heat=1 if i<6 else 2,swimming_lane=i+1 if i<6 else 3)
    with pytest.raises(ValueError,match="full"):
        move_or_swap(full,settings(),"swim",full[-1]["athlete_number"],1)


def test_existing_database_inference_preserves_override(tmp_path):
    path=tmp_path/"legacy.sqlite"
    schema=(Path(__file__).parent/"fixtures/legacy_season_schema.sql").read_text()
    with sqlite3.connect(path) as conn:
        conn.executescript(schema)
        conn.execute("INSERT INTO season_events VALUES (1,'league','League','2026-09-01',?,'x.xlsx','saved')",(LEAGUE,))
    init_season_db(path)
    assert season_snapshot(path)["events"][0]["season_year"]==2027
    with sqlite3.connect(path) as conn:
        conn.execute("UPDATE season_events SET season_year=2026")
    init_season_db(path)
    assert season_snapshot(path)["events"][0]["season_year"]==2026
    assert len(list(tmp_path.glob("*.before-seasons-*.sqlite")))==1


def entries(count, group="U/13 GIRLS", offset=0):
    from services.competition import distances,gender_for_group
    run,swim=distances(group)
    return [dict(athlete_number=str(100+offset+i),athlete_name=f"Athlete {offset+i}",group_name=group,
        gender=gender_for_group(group),run_distance=run,swim_distance=swim,
        run_entered=1,swim_entered=1,run_seed=10000+i*10,swim_seed=4000+i*10,sort_order=offset+i+1)
        for i in range(count)]


def settings(meet_type="Local",lanes=6):
    return dict(meet_type=meet_type,pool_lanes=lanes,run_positions=json.dumps(list(range(1,13))))


@pytest.mark.parametrize("discipline,prefix",[("run","running"),("swim","swimming")])
def test_manual_combine_preserves_unselected_and_rejects_other_distance(discipline,prefix):
    from services.heats import combine_selected_heats
    rows=generate_heats(entries(16),settings(),optimise=False)
    for i,row in enumerate(rows):
        row[prefix+"_heat"]=i//4+1
        row[prefix+"_lane"]=i%4+1
    programme=[{"Heat":h,"New Position":h,"Combine":False} for h in range(1,5)]
    result,_=combine_selected_heats(rows,settings(),discipline,programme,{1,3,4})
    assert [r for r in result if r["athlete_number"] in {a["athlete_number"] for a in rows[4:8]}]==rows[4:8]
    rows[-1][discipline+"_distance"]*=2
    with pytest.raises(ValueError,match="distance"):
        combine_selected_heats(rows,settings(),discipline,programme,{1,3,4})


def test_roster_save_is_explicit_atomic_and_preserves_capture(tmp_path):
    from database.repository import create_event,replace_athletes,get_event,get_athletes,save_review_roster,update_manual_time
    db=tmp_path/"event.sqlite"
    init_db(db)
    eid=create_event(db,"Review","GN","2026-09-01","SCM",6)
    replace_athletes(db,eid,generate_heats(entries(3),settings()))
    update_manual_time(db,eid,"100","run_time","02:40.00")
    before=get_athletes(db,eid)
    revision=get_event(db,eid)["heat_revision"]
    rows=[dict(r) for r in before[:2]]
    rows[0]["run_time"]="00:01.00"  # Heat review must not overwrite capture.
    with pytest.raises(ValueError,match="explicitly"):
        save_review_roster(db,eid,rows,revision)
    assert get_athletes(db,eid)==before
    with pytest.raises(ValueError,match="another session"):
        save_review_roster(db,eid,rows,revision-1,removed={"102"})
    assert get_athletes(db,eid)==before
    save_review_roster(db,eid,rows,revision,removed={"102"})
    after=get_athletes(db,eid)
    assert len(after)==2 and after[0]["run_time"]=="02:40.00"
    with sqlite3.connect(db) as conn:
        audit=conn.execute("SELECT details FROM audit_log WHERE action='EVENT_ATHLETE_REMOVED'").fetchone()
    assert json.loads(audit[0])["athlete_number"]=="102"


def test_special_needs_distances_backfill_and_seeds(tmp_path):
    from services.competition import distances
    from exporters.timedrops_json import event_for_group
    db=tmp_path/"history.sqlite"
    init_season_db(db)
    for gender in ("Male", "Female"):
        group=f"Special Needs {gender}"
        assert distances(group)==(400,50)
        assert event_for_group(group)==2
    store_event(db,NormalizedEvent("League",date(2026,9,1),LEAGUE,"x.xlsx",
        [Result("100","Athlete 0","Special Needs Female",run_time="02:00.00",swim_time="00:45.00"),
         Result("101","Athlete 1","Special Needs Male",run_distance=800,swim_distance=100)],season_year=2027))
    with sqlite3.connect(db) as conn:
        conn.execute("UPDATE season_results SET run_distance=NULL,swim_distance=NULL WHERE athlete_id=(SELECT id FROM season_athletes WHERE athlete_number='100')")
    init_season_db(db)
    init_season_db(db)
    assert len(list(tmp_path.glob("*.before-seasons-*.sqlite")))==1
    snapshot=season_snapshot(db)
    linked={a["athlete_number"]:a["id"] for a in snapshot["athletes"]}
    assert select_seed(db,linked["100"],"run",400,2027)[0]==12000
    assert select_seed(db,linked["100"],"swim",50,2027)[0]==4500
    explicit=next(r for r in snapshot["results"] if r["athlete_id"]==linked["101"])
    assert (explicit["run_distance"],explicit["swim_distance"])==(800,100)
    rows=generate_heats(entries(13,"Special Needs Female"),settings())
    assert all((r["run_distance"],r["swim_distance"])==(400,50) for r in rows)
    from services.heats import validate_heats
    for i,row in enumerate(rows):
        row.update(running_heat=1,running_lane=i+1)
    assert validate_heats(rows,settings())[0]==[]  # Manual position 13 is allowed.


def test_initial_structure_nt_and_lane_assignment():
    field=entries(13)+entries(3,"U/13 BOYS",20)
    field[0]["run_seed"]=field[0]["swim_seed"]=None
    rows=generate_heats(field,settings(),optimise=False)
    for discipline,prefix in (("run","running"),("swim","swimming")):
        heats={r[prefix+"_heat"] for r in rows}
        for heat in heats:
            members=[r for r in rows if r[prefix+"_heat"]==heat]
            assert len({r["gender"] for r in members})==1
            assert len(members)<=(15 if discipline=="run" else 6)
        assert rows[0][prefix+"_heat"]==min(r[prefix+"_heat"] for r in rows if r["gender"]=="F")
    assert centre_out(6)==[3,4,2,5,1,6]
    assert centre_out(8)==[4,5,3,6,2,7,1,8]
    small=generate_heats(entries(2),settings(lanes=8))
    assert [r["swimming_lane"] for r in small]==[4,5]
    assert [r["running_lane"] for r in small]==[2,1]
    from services.heats import reseed_run_positions
    small[0]["running_lane"],small[1]["running_lane"]=12,11
    fixed=reseed_run_positions(small)
    assert [r["running_lane"] for r in fixed]==[2,1]
    assert [(r["running_heat"],r["swimming_heat"],r["swimming_lane"]) for r in fixed]==[(r["running_heat"],r["swimming_heat"],r["swimming_lane"]) for r in small]
    assert [r["running_lane"] for r in small]==[12,11]  # No mutation before saving.


def test_local_interprovincial_and_national_underfill_rules():
    def swim_count(rows):
        return len({r["swimming_heat"] for r in rows})
    field=entries(4)+entries(2,"U/13 BOYS",10)
    assert swim_count(generate_heats(field,settings("Local")))==1
    assert swim_count(generate_heats(field,settings("Interprovincial")))==2
    assert swim_count(generate_heats(field,settings("National")))==2
    field=entries(3)+entries(3,"U/13 BOYS",10)
    assert swim_count(generate_heats(field,settings("Interprovincial")))==1
    field=entries(5)+entries(1,"U/13 BOYS",10)
    assert swim_count(generate_heats(field,settings("Local")))==1
    field=entries(6)+entries(6,"U/15 GIRLS",10)
    rows=generate_heats(field,settings("Interprovincial"))
    assert len({r["running_heat"] for r in rows})==1
    assert swim_count(rows)==2  # 50m and 100m may never mix.


def test_whole_group_mixing_fast_swim_heats_and_protected_runs():
    from services.competition import distances,category_key
    assert category_key("U/10 GIRLS") is None and distances("U/12 BOYS")== (None,None)
    rows=generate_heats(entries(8,"U/15 GIRLS")+entries(4,"U/17 GIRLS",20),settings("Interprovincial"))
    fast=[r for r in rows if r["athlete_number"] in {str(100+i) for i in range(6)}]
    assert len({r["swimming_heat"] for r in fast})==1
    heat=fast[0]["swimming_heat"]
    assert len([r for r in rows if r["swimming_heat"]==heat])==6
    rows=generate_heats(entries(4,"U/15 GIRLS")+entries(4,"U/17 GIRLS",10)+entries(2,"U/19 GIRLS",20),settings())
    assert len({r["swimming_heat"] for r in rows if r["group_name"]=="U/19 GIRLS"})==1
    for women,men,count in ((3,4,1),(6,6,1)):
        rows=generate_heats(entries(women,"MASTERS 60+ WOMEN")+entries(men,"MASTERS 70+ MEN",20),settings())
        assert len({r["running_heat"] for r in rows})==count
    rows=generate_heats(entries(8,"MASTERS 60+ WOMEN")+entries(3,"U/08 GIRLS",20),settings())
    assert len({r["running_heat"] for r in rows})==2 and len(rows)==11
    rows=generate_heats(entries(12,"MASTERS 60+ WOMEN")+entries(2,"SPECIAL NEEDS FEMALE",20),settings())
    special=next(r["running_heat"] for r in rows if r["group_name"]=="SPECIAL NEEDS FEMALE")
    assert not any(r["group_name"]=="MASTERS 60+ WOMEN" and r["running_heat"]==special for r in rows)
    assert all(sum(r["running_heat"]==h for r in rows)<=12 for h in {r["running_heat"] for r in rows})


def test_balanced_same_group_runs_are_protected_from_other_remainders():
    field=entries(16,"U/11 GIRLS")+entries(4,"U/09 GIRLS",20)
    rows=generate_heats(field,settings("Interprovincial"))
    u11=[r for r in rows if r["group_name"]=="U/11 GIRLS"]
    assert sorted(sum(r["running_heat"]==h for r in u11) for h in {r["running_heat"] for r in u11})==[8,8]
    # The four U9 runners can now join an intact U11 heat; neither U11 base
    # group donates athletes or is split up by sparse-heat repair.
    assert sorted(sum(r["running_heat"]==h for r in rows) for h in {r["running_heat"] for r in rows})==[8,12]
    field=entries(7,"U/15 GIRLS")+entries(3,"U/17 GIRLS",20)
    rows=generate_heats(field,settings("Interprovincial"))
    assert len({r["running_heat"] for r in rows})==1


@pytest.mark.parametrize("meet,count,protected",[
    ("Local",6,False),("Local",5,False),("Local",4,False),
    ("Interprovincial",5,True),("Interprovincial",4,True),("Interprovincial",3,False),
])
def test_swim_protection_and_remainder_eligibility(meet,count,protected):
    from services.heats import is_heat_protected,is_heat_eligible_for_optimisation,protection_reason
    group=entries(count,"U/09 GIRLS")
    assert is_heat_protected(group,"swim",6,meet)==protected
    assert is_heat_eligible_for_optimisation(group,"swim",6,meet)==(not protected)
    assert bool(protection_reason(group,"swim",6,meet))==protected
    rows=generate_heats(group+entries(6-count,"U/11 GIRLS",20),settings(meet))
    assert len({r["swimming_heat"] for r in rows})==(2 if protected and count<6 else 1)


@pytest.mark.parametrize("count",[5,6])
def test_interprovincial_small_runs_can_remain_intact(count):
    # Same distance alone cannot justify combining Masters and U11.
    rows=generate_heats(entries(count,"U/11 GIRLS")+entries(2,"MASTERS 60+ WOMEN",20),settings("Interprovincial"))
    assert len({r["running_heat"] for r in rows})==2
    # A weak or mixed neighbouring alternative is not a strong enough reason.
    rows=generate_heats(entries(count,"U/13 GIRLS")+entries(2,"U/19 GIRLS",20),settings("Interprovincial"))
    assert len({r["running_heat"] for r in rows})==2
    rows=generate_heats(entries(count,"U/09 GIRLS")+entries(2,"U/11 BOYS",20),settings("Interprovincial"))
    assert len({r["running_heat"] for r in rows})==2


def test_distance_national_boundaries_and_representative_compatible_pairs():
    field=entries(3,"U/08 GIRLS")+entries(3,"U/09 BOYS",20)
    rows=generate_heats(field,settings())
    assert len({r["running_heat"] for r in rows})==1
    assert len({r["swimming_heat"] for r in rows})==2  # 25m never joins 50m.
    field=entries(3,"U/13 GIRLS")+entries(3,"U/15 GIRLS",20)+entries(2,"U/13 BOYS",30)
    for meet in ("Local","Interprovincial","National"):
        rows=generate_heats(field,settings(meet))
        for prefix,discipline in (("running","run"),("swimming","swim")):
            for heat in {r[prefix+"_heat"] for r in rows}:
                members=[r for r in rows if r[prefix+"_heat"]==heat]
                assert len({r[discipline+"_distance"] for r in members})==1
                if meet=="National":
                    assert len({(r["group_name"],r["gender"]) for r in members})==1
        if meet=="Interprovincial":
            girls=[r for r in rows if r["gender"]=="F"]
            assert len({r["running_heat"] for r in girls})==1  # U13/U15 same gender first.
    # Explicit distances remain a boundary even for otherwise identical categories.
    field=entries(2)
    field[1].update(run_distance=400,swim_distance=100)
    rows=generate_heats(field,settings())
    assert len({r["running_heat"] for r in rows})==len({r["swimming_heat"] for r in rows})==2


def test_league_heat_count_precedes_former_destination_preferences():
    field=entries(4,"U/13 GIRLS")+entries(3,"U/15 BOYS",20)+entries(6,"U/17 GIRLS",30)
    rows=generate_heats(field,settings())
    assert len({r["running_heat"] for r in rows})==1


@pytest.mark.parametrize("discipline,prefix",[("run","running"),("swim","swimming")])
def test_special_needs_uses_same_distance_gender_and_seeds(discipline,prefix):
    field=entries(2,"SPECIAL NEEDS FEMALE")+entries(4,"MASTERS 60+ WOMEN",20)+entries(4,"U/09 GIRLS",30)
    for row in field:
        row[discipline+"_seed"]=20000 if row["group_name"]=="MASTERS 60+ WOMEN" else 10000
    rows=generate_heats(field,settings())
    target=next(r[prefix+"_heat"] for r in rows if r["group_name"]=="SPECIAL NEEDS FEMALE")
    assert {r["group_name"] for r in rows if r[prefix+"_heat"]==target}=={"SPECIAL NEEDS FEMALE","U/09 GIRLS"}
    # Changing only seed suitability selects the older Masters instead.
    for row in field:
        row[discipline+"_seed"]=20000 if row["group_name"]=="U/09 GIRLS" else 10000
    rows=generate_heats(field,settings())
    target=next(r[prefix+"_heat"] for r in rows if r["group_name"]=="SPECIAL NEEDS FEMALE")
    assert {r["group_name"] for r in rows if r[prefix+"_heat"]==target}=={"SPECIAL NEEDS FEMALE","MASTERS 60+ WOMEN"}
    # Running favours gender; swimming favours seed coherence within valid options.
    field=entries(2,"SPECIAL NEEDS FEMALE")+entries(4,"MASTERS 60+ MEN",20)+entries(4,"U/09 GIRLS",30)
    for row in field:
        row[discipline+"_seed"]=20000 if row["group_name"]=="U/09 GIRLS" else 10000
    rows=generate_heats(field,settings())
    target=next(r[prefix+"_heat"] for r in rows if r["group_name"]=="SPECIAL NEEDS FEMALE")
    assert {r["group_name"] for r in rows if r[prefix+"_heat"]==target}=={"SPECIAL NEEDS FEMALE", "U/09 GIRLS" if discipline=="run" else "MASTERS 60+ MEN"}


def test_adult_cluster_and_u19_bridge_in_both_disciplines():
    field=entries(2,"SENIOR WOMEN")+entries(2,"MASTERS 40+ WOMEN",20)+entries(2,"U/19 GIRLS",30)
    rows=generate_heats(field,settings())
    assert len({r["running_heat"] for r in rows})==len({r["swimming_heat"] for r in rows})==1


def test_generation_ignores_previous_assignments_and_input_order():
    field=entries(8,"U/15 GIRLS")+entries(4,"U/17 GIRLS",20)+entries(2,"U/19 BOYS",30)
    field[0]["run_seed"]=field[0]["swim_seed"]=None
    expected=generate_heats(field,settings())
    for row in field:
        row.update(running_heat=99,running_lane=18,swimming_heat=25,swimming_lane=1)
    assert sorted(generate_heats(list(reversed(field)),settings()),key=lambda r:r["athlete_number"])==sorted(expected,key=lambda r:r["athlete_number"])


def test_manual_distance_exception_keeps_automatic_boundary():
    from services.heats import move_or_swap,validate_heats
    field=generate_heats(entries(1,"U/11 GIRLS")+entries(1,"U/13 GIRLS",20),settings())
    assert len({r["running_heat"] for r in field})==2
    moved=move_or_swap(field,settings(),"run",field[1]["athlete_number"],field[0]["running_heat"])
    errors,warnings=validate_heats(moved,settings())
    assert not errors and any("different distances" in warning for warning in warnings)
    from services.heat_outputs import build_package
    event=dict(settings(),name="Manual exception",host_team="GN",start_date="2026-09-01",
               course="SCM",season_year=2027,heat_revision=1)
    assert len(build_package(moved,event,{}))==5  # Mixed run distances do not break swim outputs.
    moved[1]["swim_distance"]=100
    with pytest.raises(ValueError,match="different TimeDrops events"):
        build_package(moved,event,{})  # The existing TimeDrops format guard remains.


def test_manual_programme_compacts_heats_and_preserves_roster_and_capture():
    from copy import deepcopy
    from services.heats import apply_heat_programme
    previous=generate_heats(entries(13),settings(),optimise=False)
    for row in previous:
        row.update(run_position=3,run_time="02:10.00")
    edited=deepcopy(previous)
    for row in edited:
        row["running_heat"]=9
    programme=[{"Heat":1,"Delete":True},{"Heat":2,"Delete":True},{"Heat":9},{"Heat":10}]
    saved=apply_heat_programme(edited,previous,settings(),"run",programme)
    assert {r["running_heat"] for r in saved}=={1}
    assert {r["running_lane"] for r in saved}==set(range(1,14))  # Manual capacity override remains.
    assert sorted(saved,key=lambda r:r["run_seed"])[0]["running_lane"]==13
    assert all(r["run_position"]==3 and r["run_time"]=="02:10.00" for r in saved)
    assert [(r["swimming_heat"],r["swimming_lane"]) for r in saved]==[(r["swimming_heat"],r["swimming_lane"]) for r in previous]
    assert {r["running_heat"] for r in edited}=={9}  # No mutation before successful save.
    with pytest.raises(ValueError,match="Move all athletes out"):
        apply_heat_programme(previous,previous,settings(),"run",programme)


def test_manual_swim_batch_moves_reseed_and_preserve_explicit_lanes_in_exports():
    from copy import deepcopy
    from io import BytesIO
    from services.heats import apply_heat_programme
    from services.heat_outputs import build_package
    from parsers.master_entries import parse_master_entries
    previous=generate_heats(entries(8),settings(),optimise=False)
    edited=deepcopy(previous)
    moving=[r for r in edited if r["swimming_heat"]==2][:2]
    for row in moving:
        row["swimming_heat"]=1
    moving[0]["swimming_lane"]=1
    override=moving[0]["athlete_number"]
    saved=apply_heat_programme(edited,previous,settings(),"swim",[{"Heat":2},{"Heat":1},{"Heat":3}],{override})
    assert next(r for r in saved if r["athlete_number"]==override)["swimming_lane"]==1
    assert all(len({r["swimming_lane"] for r in saved if r["swimming_heat"]==h})==4 for h in (1,2))
    assert [(r["running_heat"],r["running_lane"]) for r in saved]==[(r["running_heat"],r["running_lane"]) for r in previous]
    event=dict(settings(),name="Edited programme",host_team="GN",start_date="2026-09-01",course="SCM",season_year=2027,heat_revision=1)
    files=build_package(saved,event,{})
    parsed=parse_master_entries(BytesIO(files["Master Entries Heats.xlsx"][1]))["athletes"]
    assignments=lambda rows:{r["athlete_number"]:(r["running_heat"],r["running_lane"],r["swimming_heat"],r["swimming_lane"]) for r in rows}
    assert assignments(parsed)==assignments(saved)
    races=json.loads(files["meet_program.json"][1])["meetSessions"][0]["sessionRaces"]
    exported={lane["laneSwimmerId"]:(race["raceHeatNumber"],lane["laneNumber"]) for race in races for lane in race["raceLanes"]}
    assert exported=={r["athlete_number"]:(r["swimming_heat"],r["swimming_lane"]) for r in saved}
    for row in edited:
        row["swimming_heat"]=1
    with pytest.raises(ValueError,match="exceeds the pool capacity"):
        apply_heat_programme(edited,previous,settings(),"swim",[{"Heat":1}])


def test_print_layout_and_simplified_export_headers():
    from io import BytesIO
    import pdfplumber
    import openpyxl
    from exporters.heats import build_heats_pdf,build_heats_xlsx
    from parsers.master_entries_pdf import parse_master_entries_pdf
    event=dict(settings("National"),name="GNB League 2",start_date="2026-09-01",season_year=2027,
        host_team="GN",course="SCM",heat_revision=1)
    rows=generate_heats(entries(24),event)
    pdf=build_heats_pdf(rows,event).getvalue()
    with pdfplumber.open(BytesIO(pdf)) as document:
        text=document.pages[0].extract_text()
        assert "Heat 1 -" in text and "Heat 2 -" in text
        assert all(r["athlete_name"] in text for r in rows)
        assert not any(label in text for label in ("Run positions:","Meet type:","Course:","Host:"))
        assert any(c["size"]==8.5 for c in document.pages[0].chars)
    parsed=parse_master_entries_pdf(BytesIO(pdf))
    assert len(parsed["athletes"])==24
    assert {r["athlete_number"]:(r["running_heat"],r["running_lane"],r["swimming_heat"],r["swimming_lane"]) for r in parsed["athletes"]}=={
        r["athlete_number"]:(r["running_heat"],r["running_lane"],r["swimming_heat"],r["swimming_lane"]) for r in rows}
    ws=openpyxl.load_workbook(build_heats_xlsx(rows,event)).active
    text=" ".join(str(row[0].value or "") for row in ws)
    assert all(label not in text for label in ("Pool lanes:","Run positions:","Meet type:","Course:","Host:"))


def test_manual_assignments_reach_every_export_and_phase2_invalidates(tmp_path):
    from io import BytesIO
    import openpyxl
    import pytest
    from database.repository import create_event,replace_athletes,get_event,get_athletes,save_heat_assignments,approve_heats,assign_athlete_to_run_heat
    from services.heat_outputs import generate_outputs,current_outputs
    from parsers.master_entries import parse_master_entries
    from parsers.master_entries_pdf import parse_master_entries_pdf
    db=tmp_path/"events.sqlite"
    init_db(db)
    eid=create_event(db,"Acceptance Meet","GN","2026-09-01","SCM",6,heat_source="generated")
    replace_athletes(db,eid,entries(9))
    event=dict(get_event(db,eid))
    rows=generate_heats(get_athletes(db,eid),event)
    rows[0].update(running_heat=4,running_lane=12,swimming_heat=4,swimming_lane=3)
    save_heat_assignments(db,eid,rows,event["heat_revision"])
    approve_heats(db,eid,get_event(db,eid)["heat_revision"])
    files=generate_outputs(db,eid,{})
    assert set(files)=={"Combined Heats.pdf","Master Entries Heats.xlsx","Athlete Run Swim Lane Sheet.xlsx","Swim Timekeeper Sheets.xlsx","meet_program.json"}
    for parsed in (parse_master_entries(BytesIO(files["Master Entries Heats.xlsx"][1])),
                   parse_master_entries_pdf(BytesIO(files["Combined Heats.pdf"][1]))):
        moved=next(r for r in parsed["athletes"] if r["athlete_number"]=="100")
        assert (moved["running_heat"],moved["running_lane"],moved["swimming_heat"],moved["swimming_lane"])==(4,12,4,3)
        assert moved["athlete_name"]=="Athlete 0"
        assert parsed["metadata"]["season_year"]==2027
    program=json.loads(files["meet_program.json"][1])
    race=next(r for r in program["meetSessions"][0]["sessionRaces"] if r["raceHeatNumber"]==4)
    assert race["raceLanes"][0]["laneSwimmerId"]=="100" and race["raceLanes"][0]["laneNumber"]==3
    assert race["raceLanes"][0]["laneSeedTime"]==0
    workbook=openpyxl.load_workbook(BytesIO(files["Athlete Run Swim Lane Sheet.xlsx"][1]))
    mapped=next(row for row in workbook.active.iter_rows(min_row=2,values_only=True) if row[0]=="100")
    assert mapped[3:]==(4,4,3)
    workbook=openpyxl.load_workbook(BytesIO(files["Swim Timekeeper Sheets.xlsx"][1]))
    assert all("2026-09-01" in sheet["A1"].value for sheet in workbook)
    timed=next(row for row in workbook["Lane3"].iter_rows(min_row=8,values_only=True) if row[1]=="100")
    assert timed[4]==4
    assert current_outputs(db,eid)
    from services.heats import reorder_heat
    before=get_athletes(db,eid)
    reordered=reorder_heat(before,"run",4,1)
    reordered=reorder_heat(reordered,"swim",4,1)
    assert [(r["running_lane"],r["swimming_lane"]) for r in reordered]==[(r["running_lane"],r["swimming_lane"]) for r in before]
    save_heat_assignments(db,eid,reordered,get_event(db,eid)["heat_revision"])
    assert get_event(db,eid)["heat_status"]=="stale" and not current_outputs(db,eid)
    approve_heats(db,eid,get_event(db,eid)["heat_revision"])
    reordered_files=generate_outputs(db,eid,{})
    parsed=parse_master_entries(BytesIO(reordered_files["Master Entries Heats.xlsx"][1]))
    first=next(r for r in parsed["athletes"] if r["athlete_number"]=="100")
    assert (first["running_heat"],first["swimming_heat"])==(1,1)
    assert json.loads(reordered_files["meet_program.json"][1])["meetSessions"][0]["sessionRaces"][0]["raceLanes"][0]["laneSwimmerId"]=="100"
    assign_athlete_to_run_heat(db,eid,"100",5)
    assert get_event(db,eid)["heat_status"]=="stale" and not current_outputs(db,eid)
    with pytest.raises(ValueError,match="approve"):
        generate_outputs(db,eid,{})


def test_captured_results_share_seed_and_report_history_without_erasing_published_points(tmp_path):
    import pytest
    from database.repository import create_event,replace_athletes,update_manual_time
    from services.workflow_history import save_workflow_results
    workflow=tmp_path/"workflow.sqlite"
    history=tmp_path/"history.sqlite"
    init_db(workflow)
    init_season_db(history)
    eid=create_event(workflow,"League","GN","2026-09-01","SCM",6)
    rows=entries(1)
    replace_athletes(workflow,eid,rows)
    update_manual_time(workflow,eid,"100","run_time","02:40.00")
    update_manual_time(workflow,eid,"100","swim_time","00:40.00")
    hid=save_workflow_results(workflow,history,eid)
    assert save_workflow_results(workflow,history,eid)==hid
    snapshot=season_snapshot(history,2027)
    assert len(snapshot["events"])==len(snapshot["results"])==1
    result=snapshot["results"][0]
    assert result["total_points"]==""
    assert select_seed(history,result["athlete_id"],"run",800,2027)[0]==16000
    published=NormalizedEvent("League",date(2026,9,1),LEAGUE,"official.xlsx",
        [Result("100","Athlete 0","U/13 GIRLS",run_time="02:39.00",swim_time="00:39.00",total_points="2100")],season_year=2027)
    assert store_event(history,published,overwrite=True)==hid
    with pytest.raises(ValueError,match="Published results"):
        save_workflow_results(workflow,history,eid)
    assert season_snapshot(history,2027)["results"][0]["total_points"]=="2100"
