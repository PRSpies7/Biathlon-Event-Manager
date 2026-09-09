"""Generate a filterable season workbook using the existing openpyxl exporters' pattern."""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.season_reports import age_group_sort_key, athlete_result_rows, report_sort_key

ATHLETE_HEADERS = ["Athlete number", "Athlete name", "School", "Age group", "Province", "Team", "Affiliated", "Events attended", "Event awards",
                   "Run time", "Swim time", "Total points", "Result event"]
RESULT_COLUMNS = ["event_id", "event_name", "event_date", "competition_type", "athlete_number", "athlete_name",
                  "category", "position", "run_time", "running_points", "swim_time", "swimming_points",
                  "bonus_points", "total_points", "status", "school", "province", "team", "annotations"]
EVENT_COLUMNS = ["id", "name", "event_date", "competition_type", "source_filename", "imported_at", "result_count"]
AWARD_COLUMNS = ["event_id", "event_name", "event_date", "athlete_number", "athlete_name", "award_type", "placing"]


def _value(key, value):
    if key == "event_date" and value:
        return date.fromisoformat(value)
    if key == "imported_at" and value:
        return datetime.fromisoformat(value).replace(tzinfo=None)  # labeled UTC in the sheet
    if key.endswith("points") or key == "position":
        try:
            number = Decimal(value)
            return float(number) if number.is_finite() else value
        except (InvalidOperation, TypeError):
            return value
    return value


def build_season_results_xlsx(snapshot, *, athlete_report=None, include_points=True) -> BytesIO:
    athletes = athlete_result_rows(snapshot) if athlete_report is None else sorted(athlete_report, key=report_sort_key)
    headers = list(athletes[0]) if athletes else ATHLETE_HEADERS
    if not include_points:
        headers = [h for h in headers if "points" not in h.casefold()]
        athletes = sorted(athletes, key=lambda r: (age_group_sort_key(r.get("Age group", "")), r["Athlete name"].casefold()))
    datasets = [("Athletes", headers, [[row.get(h, "") for h in headers] for row in athletes])]
    for name, key, columns in (("Event Results", "results", RESULT_COLUMNS),
                               ("Imported Events", "events", EVENT_COLUMNS),
                               ("Event Awards", "awards", AWARD_COLUMNS)):
        if not include_points:
            columns = [c for c in columns if not c.endswith("points")]
        labels = ["Age group" if c == "category" else "Imported at (UTC)" if c == "imported_at" else c.replace("_", " ").capitalize() for c in columns]
        source_rows = sorted(snapshot[key], key=report_sort_key) if key == "results" else snapshot[key]
        datasets.append((name, labels, [[_value(c, row[c]) for c in columns] for row in source_rows]))
    if include_points and snapshot.get("records"):
        columns = ["category", "athlete_number", "athlete_name", "total_points", "source_filename", "imported_at"]
        labels = ["Category", "Athlete number", "Athlete name", "Total points", "Source filename", "Imported at (UTC)"]
        datasets.append(("Current Records", labels, [[_value(c, row[c]) for c in columns] for row in sorted(snapshot["records"], key=report_sort_key)]))
    return _build_workbook(datasets)


def build_report_xlsx(sheet_name, headers, rows) -> BytesIO:
    return _build_workbook([(sheet_name, headers, [[row.get(h, "") for h in headers] for row in rows])])


def build_qualified_athletes_xlsx(snapshot, rows) -> BytesIO:
    from services.season_reports import qualified_athlete_report
    headers = qualified_athlete_report(snapshot)[0]
    return with_insights(build_report_xlsx("Qualified Athletes", headers, rows), snapshot)


def build_top_athletes_xlsx(snapshot):
    from services.season_reports import top_athlete_report
    headers, rows = top_athlete_report(snapshot)
    stream = with_insights(build_report_xlsx("Top Athletes", headers, rows), snapshot)
    workbook = load_workbook(stream)
    detail = load_workbook(build_season_results_xlsx(snapshot))["Event Results"]
    sheet = workbook.create_sheet("Event Details")
    from copy import copy
    for row in detail:
        for cell in row:
            dest = sheet.cell(cell.row, cell.column, cell.value)
            dest._style = copy(cell._style)
            dest.data_type = cell.data_type
    for key, dimension in detail.column_dimensions.items():
        sheet.column_dimensions[key] = copy(dimension)
    for key, dimension in detail.row_dimensions.items():
        sheet.row_dimensions[key] = copy(dimension)
    sheet.freeze_panes = detail.freeze_panes
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def with_insights(stream, snapshot, *, affiliation_only=False):
    from services.season_insights import insight_sections
    workbook = load_workbook(stream)
    sheet = workbook.create_sheet("Insights", 1)
    for title, headers, rows in insight_sections(snapshot, affiliation_only=affiliation_only):
        sheet.append([f"{title} — {len(rows)}" if "Affiliation" in title else title])
        sheet.cell(sheet.max_row, 1).font = Font(bold=True, color="17365D", size=12)
        sheet.merge_cells(start_row=sheet.max_row, start_column=1, end_row=sheet.max_row, end_column=8)
        sheet.append(headers)
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.font = Font(bold=True, color="FFFFFF")
        for row in rows:
            sheet.append([row.get(h, "") for h in headers])
        if not rows:
            sheet.append(["No matching athletes or events."])
        if rows and title in {"Participation by event", "Age groups with fewer than six athletes (latest category)"}:
            chart = BarChart()
            chart.title = title
            col = headers.index("Athletes") + 1
            chart.add_data(Reference(sheet, min_col=col, min_row=header_row, max_row=sheet.max_row), titles_from_data=True)
            chart.set_categories(Reference(sheet, min_col=1, min_row=header_row + 1, max_row=sheet.max_row))
            chart.height = 8
            chart.width = 20
            sheet.add_chart(chart, f"J{header_row}")
            for _ in range(max(0, 17 - len(rows))):
                sheet.append([""])
        sheet.append([""])
        sheet.append([""])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[row[0].row].height = 45
    for i in range(1, 9):
        sheet.column_dimensions[get_column_letter(i)].width = 28
    sheet.freeze_panes = "A3"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_qualified_schools_xlsx(summary, athlete_headers, athletes) -> BytesIO:
    from services.school_reports import SUMMARY_HEADERS
    return _build_workbook([
        ("Qualified Schools", SUMMARY_HEADERS, [[r.get(h, "") for h in SUMMARY_HEADERS] for r in summary]),
        ("School Athlete Scores", athlete_headers, [[r.get(h, "") for h in athlete_headers] for r in athletes]),
    ])


def _build_workbook(datasets) -> BytesIO:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, headers, rows in datasets:
        season_matrix = name in {"Top Athletes", "Qualified Athletes"}
        school_matrix = name == "School Athlete Scores"
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "H2" if school_matrix else "E2" if season_matrix else "C2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.data_type = "s"
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 75 if school_matrix else 60 if season_matrix else 32
        for i, label in enumerate(headers, 1):
            width = 65 if label in {"Event awards", "Missing to qualify"} else 34 if any(t in label.casefold() for t in ("name", "school", "source")) else 23
            if season_matrix and i > 10:
                width = 30
            sheet.column_dimensions[get_column_letter(i)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                # Uploaded names are literal text, including leading '='; never Excel formulas.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
                elif "points" in headers[cell.column - 1].casefold() or (season_matrix and cell.column > 10):
                    cell.number_format = "0.00"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            longest = max((len(str(c.value or "")) / sheet.column_dimensions[c.column_letter].width for c in row), default=1)
            sheet.row_dimensions[row[0].row].height = max(30, 15 * (int(longest) + 1))
            if name in {"Qualified Schools", "Qualified Athletes"}:
                status = row[headers.index("Qualified")]
                qualified = status.value == "Yes"
                status.fill = PatternFill("solid", fgColor="C6EFCE" if qualified else "FFC7CE")
                status.font = Font(bold=True, color="006100" if qualified else "9C0006")
                if name == "Qualified Athletes":
                    athlete = row[headers.index("Athlete name")]
                    athlete.fill = PatternFill("solid", fgColor="C6EFCE" if qualified else "FFC7CE")
                    athlete.font = Font(bold=True, color="006100" if qualified else "9C0006")
                if qualified:
                    school = row[headers.index("School")]
                    school.fill = PatternFill("solid", fgColor="C6EFCE")
                    school.font = Font(bold=True, color="006100")
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
