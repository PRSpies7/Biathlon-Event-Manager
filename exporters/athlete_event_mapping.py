from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def build_printable_athlete_mapping_xlsx(athletes: list[dict[str, Any]]) -> BytesIO:
    """Build the printable Phase 1 athlete-to-event mapping workbook."""
    rows = sorted(
        athletes,
        key=lambda athlete: " ".join(str(athlete.get("athlete_name") or "").split()).casefold(),
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Athlete Event Mapping"
    ws.append([
        "Athlete Number",
        "Athlete Name",
        "Age Group",
        "Run Heat",
        "Swim Heat",
        "Swim Lane",
    ])

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    for athlete in rows:
        ws.append([
            str(athlete.get("athlete_number") or ""),
            str(athlete.get("athlete_name") or ""),
            str(athlete.get("group_name") or ""),
            athlete.get("running_heat") if athlete.get("running_heat") is not None else "",
            athlete.get("swimming_heat") if athlete.get("swimming_heat") is not None else "",
            athlete.get("swimming_lane") if athlete.get("swimming_lane") is not None else "",
        ])

    for row_number in range(2, ws.max_row + 1):
        ws.row_dimensions[row_number].height = 19
        for cell in ws[row_number]:
            cell.alignment = Alignment(vertical="center")
        # Athlete numbers are identifiers, including when they contain leading zeroes.
        ws.cell(row=row_number, column=1).number_format = "@"

    for column, width in {
        "A": 18,
        "B": 32,
        "C": 18,
        "D": 12,
        "E": 12,
        "F": 12,
    }.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.print_title_rows = "1:1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2
    ws.print_area = f"A1:F{ws.max_row}"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
