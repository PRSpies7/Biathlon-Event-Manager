"""Update a copy of the supplied records workbook, preserving its existing layout."""
from datetime import date, time
from decimal import Decimal
from io import BytesIO
import re

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from parsers.season_records import record_category_key as category_key, records_template_rows
from services.season_reports import numeric_points


def excel_time(value):
    if not re.fullmatch(r"\d+(?::\d{2}){1,2}(?:\.\d+)?", value):
        return value
    parts = value.split(":")
    seconds = Decimal(parts[-1])
    return time(int(parts[0]) if len(parts) == 3 else 0, int(parts[-2]), int(seconds), int((seconds % 1) * 1000000))


def updated_records_xlsx(snapshot):
    template = snapshot.get("record_templates", [])
    if not template:
        raise ValueError("Upload and save your formatted records workbook first.")
    workbook = load_workbook(BytesIO(template[0]["content"]))
    layout = records_template_rows(workbook)
    if not layout:
        raise ValueError("The saved workbook is a simple reference list. Upload the formatted GN records workbook to download an updated template.")
    results = {}
    for r in sorted(snapshot["results"], key=lambda r: (r["event_date"], r["event_id"], r["athlete_number"])):
        points = numeric_points(r["total_points"])
        if points is None:
            continue
        key = category_key(r["category"])
        if key not in results or points > numeric_points(results[key]["total_points"]):
            results[key] = r  # Earliest occurrence wins an equal season best.
    changes = []
    for sheet, row in layout:
        result = results.get(category_key(sheet.cell(row, 2).value))
        baseline = numeric_points(sheet.cell(row, 5).value)
        if not result or baseline is None or numeric_points(result["total_points"]).quantize(Decimal(".01")) <= baseline.quantize(Decimal(".01")):
            continue
        points = float(numeric_points(result["total_points"]))
        old_name = sheet.cell(row, 3).value
        values = {3: result["athlete_name"], 4: float(baseline), 5: points, 6: "(GN)",
                  7: date.fromisoformat(result["event_date"]).year, 10: result["event_name"],
                  11: excel_time(result["run_time"]), 12: excel_time(result["swim_time"])}
        for column, value in values.items():
            cell = sheet.cell(row, column, value)
            if isinstance(value, str):
                cell.data_type = "s"
        changes.append([sheet.cell(row, 2).value, result["athlete_number"], result["athlete_name"], old_name,
                        float(baseline), points, result["event_name"], date.fromisoformat(result["event_date"]),
                        result["run_time"], result["swim_time"]])
    if snapshot["events"]:
        year = max(e["event_date"] for e in snapshot["events"])[:4]
        for sheet in workbook:
            if isinstance(sheet["B2"].value, str) and "BIATHLON RECORDS" in sheet["B2"].value.upper():
                sheet["B2"] = re.sub(r"\b20\d{2}\b", year, sheet["B2"].value)
    sheet = workbook.create_sheet("New Records")
    sheet.append(["Age group", "Athlete number", "Athlete name", "Previous holder", "Previous record points", "New record points", "Event", "Event date", "Run time", "Swim time"])
    for row in changes:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
    from openpyxl.utils import get_column_letter
    for i in range(1, 11):
        sheet.column_dimensions[get_column_letter(i)].width = 28
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
        row[7].number_format = "yyyy-mm-dd"
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
