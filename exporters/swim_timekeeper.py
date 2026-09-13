from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

BLACK = "000000"


def build_swim_timekeeper_xlsx(athletes: list[dict[str, Any]], event_name: str, pool_lanes: int, event_date: str | None = None) -> BytesIO:
    wb = Workbook()
    wb.remove(wb.active)

    thin = Side(style="thin", color="A6A6A6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for lane in range(1, int(pool_lanes) + 1):
        ws = wb.create_sheet(f"Lane{lane}")
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A8"

        ws.merge_cells("A1:F1")
        ws["A1"] = f"{event_name} · {event_date}" if event_date else event_name
        ws["A1"].font = Font(bold=True, size=16, color=BLACK)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 42 if event_date else 26

        ws.merge_cells("A2:F2")
        ws["A2"] = f"Swim Timekeeper - Lane {lane}"
        ws["A2"].font = Font(bold=True, size=12, color=BLACK)
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        ws.merge_cells("A4:F4")
        ws["A4"] = "Timekeeper Name: _________________________________________________"
        ws["A4"].font = Font(bold=True, size=11, color=BLACK)
        ws["A4"].alignment = Alignment(vertical="center")
        ws.row_dimensions[4].height = 22

        headers = ["Lane", "#", "Athlete Name", "Age group", "Heat", "Time"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(7, col, header)
            cell.font = Font(bold=True, color=BLACK)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        ws.row_dimensions[7].height = 22

        lane_athletes = [a for a in athletes if a.get("swimming_lane") == lane]
        heat_numbers = sorted({int(a["swimming_heat"]) for a in athletes if a.get("swimming_heat") is not None})
        by_heat = {h: [] for h in heat_numbers}
        for a in lane_athletes:
            by_heat.setdefault(int(a["swimming_heat"]), []).append(a)

        # One row per heat, even when the lane is empty, matching the physical timing workflow.
        row = 8
        for heat in heat_numbers:
            entries = sorted(by_heat.get(heat, []), key=lambda a: a.get("sort_order", 0))
            if not entries:
                entries = [None]
            for a in entries:
                values = [
                    lane,
                    a.get("athlete_number") if a else None,
                    a.get("athlete_name") if a else None,
                    a.get("group_name") if a else None,
                    heat,
                    None,
                ]
                for col, value in enumerate(values, 1):
                    cell = ws.cell(row, col, value)
                    cell.border = border
                    cell.alignment = Alignment(vertical="center", horizontal="center" if col in (1, 2, 5, 6) else "left", wrap_text=True)
                    if col == 6:
                        cell.font = Font(bold=True, size=12)
                # Excel stores row heights in points. Apply the required value
                # to every generated heat row, including empty lane/heat rows.
                ws.row_dimensions[row].height = 30
                row += 1

        last_row = max(row - 1, 7)
        ws.auto_filter.ref = f"A7:F{last_row}"
        widths = {"A": 9, "B": 12, "C": 34, "D": 24, "E": 10, "F": 22}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        # Print-ready: fit to one page wide and repeat the title/timekeeper/header block if the lane spills over.
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True, autoPageBreaks=False)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:7"
        ws.print_area = f"A1:F{last_row}"
        ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.35, bottom=0.35, header=0.15, footer=0.15)
        ws.oddFooter.center.text = "Page &P of &N"
        ws.oddFooter.center.size = 8
        ws.oddFooter.center.font = "Arial"

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
