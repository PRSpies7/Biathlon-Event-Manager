from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


def build_master_import_xlsx(athletes: list[dict[str, Any]]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "FOR IMPORT"
    headers = ["Athlete nr", "Athlete Name", "runtime", "swimtime"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="left")
        cell.number_format = "@"

    for a in athletes:
        # The supplied reference output contains only athletes with both valid times.
        if not a.get("run_time") or not a.get("swim_time"):
            continue
        ws.append([
            str(a["athlete_number"]),
            str(a["athlete_name"]),
            str(a["run_time"]),
            str(a["swim_time"]),
        ])
        for c in ws[ws.max_row]:
            c.number_format = "@"

    widths = {"A": 14, "B": 34, "C": 16, "D": 16}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
