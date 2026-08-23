from __future__ import annotations

import re
from typing import Any

import openpyxl

TIME_RE = re.compile(r"^(?:\d+:)?\d{2}:\d{2}\.\d{2}$")
HEAT_RE = re.compile(r"^Heat\s*(.+)$", re.IGNORECASE)


def parse_heat_numbers(label: str) -> tuple[int, ...]:
    nums = [int(x) for x in re.findall(r"\d+", label)]
    return tuple(dict.fromkeys(nums))


def parse_run_results_excel(path_or_file) -> dict[str, Any]:
    wb = openpyxl.load_workbook(path_or_file, read_only=True, data_only=True)
    if "Run times" not in wb.sheetnames:
        raise ValueError("Run Results workbook must contain a 'Run times' worksheet.")
    ws = wb["Run times"]

    blocks: list[dict[str, Any]] = []
    current = None

    for row_num, row in enumerate(ws.iter_rows(values_only=True), start=1):
        first = row[0] if len(row) > 0 else None
        second = row[1] if len(row) > 1 else None
        s = str(first).strip() if first is not None else ""
        m = HEAT_RE.match(s)
        if m:
            if current is not None:
                blocks.append(current)
            label = m.group(1).strip()
            current = {
                "label": f"Heat {label}",
                "heat_numbers": parse_heat_numbers(label),
                "results": [],
                "start_row": row_num,
            }
            continue
        if current is None:
            continue
        if str(first).strip().lower() == "lap":
            continue
        try:
            position = int(first)
        except (TypeError, ValueError):
            continue
        if not isinstance(second, str):
            continue
        runtime = second.strip()
        if not TIME_RE.match(runtime):
            continue
        current["results"].append({"position": position, "run_time": runtime, "row": row_num})

    if current is not None:
        blocks.append(current)

    if not blocks:
        raise ValueError("No run result heat sections were found.")

    # Later duplicate sections win, but are reported to the caller.
    by_key: dict[tuple[int, ...], dict[str, Any]] = {}
    revisions: list[dict[str, Any]] = []
    for block in blocks:
        key = tuple(block["heat_numbers"])
        if key in by_key:
            revisions.append({"heat_numbers": key, "replaced_start_row": by_key[key]["start_row"], "replacement_start_row": block["start_row"]})
        by_key[key] = block

    final_blocks = list(by_key.values())
    source_timed_count = sum(len(b.get("results", [])) for b in final_blocks)
    return {
        "blocks": final_blocks,
        "revisions": revisions,
        "imported_heat_numbers": sorted({n for b in final_blocks for n in b["heat_numbers"]}),
        "source_timed_count": source_timed_count,
    }
