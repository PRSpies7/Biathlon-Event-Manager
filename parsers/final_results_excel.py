from __future__ import annotations

from datetime import time, timedelta
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from validation.validators import normalize_time, validate_master, validate_times


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _time(value: Any) -> str | None:
    if value is None or _text(value) in {"", "0", "0.00", "00:00.00", "0:00.00", "00.00"}:
        return None
    if isinstance(value, time):
        seconds = value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1_000_000
    elif isinstance(value, timedelta):
        seconds = value.total_seconds()
    else:
        normalized = normalize_time(value)
        if normalized is None:
            raise ValueError("Expected a time such as 04:12.50, or a blank/00:00.00 for a missing result.")
        return normalized
    hundredths = round(seconds * 100)
    if hundredths < 0:
        raise ValueError("Times cannot be negative.")
    minutes, remainder = divmod(hundredths, 6000)
    return None if hundredths == 0 else f"{minutes:02d}:{remainder // 100:02d}.{remainder % 100:02d}"


def parse_final_results_excel(source) -> list[dict[str, Any]]:
    """Read the exported final-results workbook without changing saved events.

    Formulas are rejected so stale cached results cannot silently be exported.
    """
    try:
        workbook = load_workbook(source, read_only=True, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise ValueError("Could not read this Excel workbook. Upload a valid .xlsx final-results file.") from exc
    try:
        sheet = workbook["FOR IMPORT"] if "FOR IMPORT" in workbook.sheetnames else workbook.worksheets[0]
        rows = iter(sheet.iter_rows())
        header = next(rows, ())
        columns = {_text(cell.value).casefold(): index for index, cell in enumerate(header)}
        required = {"Athlete nr": "athlete_number", "Athlete Name": "athlete_name",
                    "runtime": "run_time", "swimtime": "swim_time"}
        missing = [name for name in required if name.casefold() not in columns]
        if missing:
            raise ValueError("Missing columns: " + ", ".join(missing) + ". Use the final-results Excel export with its original headers.")
        athletes = []
        errors = []
        for row_number, cells in enumerate(rows, start=2):
            if all(cell.value is None or _text(cell.value) == "" for cell in cells):
                continue
            athlete = {}
            fields = {**required, "athleteSurname": "athlete_surname"}
            for label, field in fields.items():
                index = columns.get(label.casefold())
                cell = cells[index] if index is not None and index < len(cells) else None
                value = cell.value if cell is not None else None
                if cell is not None and cell.data_type == "f":
                    errors.append(f"Row {row_number}, {label}: replace the formula with its calculated value before uploading.")
                    value = None
                elif cell is not None and cell.data_type == "e":
                    errors.append(f"Row {row_number}, {label}: correct the Excel error {value}.")
                    value = None
                try:
                    athlete[field] = _time(value) if field in {"run_time", "swim_time"} else _text(value)
                except ValueError as exc:
                    errors.append(f"Row {row_number}, {label}: {exc}")
                    athlete[field] = None
            athletes.append(athlete)
        errors.extend(validate_master(athletes))
        errors.extend(validate_times(athletes))
        if errors:
            raise ValueError("Please correct the workbook before converting:\n" + "\n".join(errors[:50]))
        if not any(a["run_time"] or a["swim_time"] for a in athletes):
            raise ValueError("The workbook contains no results to convert.")
        return athletes
    finally:
        workbook.close()
