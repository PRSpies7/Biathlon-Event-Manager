from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import openpyxl

from .master_entries_pdf import parse_master_entries_pdf

HEAT_RE = re.compile(r"^Heat\s*(\d+)\s*-\s*(.*)$", re.IGNORECASE)
INTERPROVINCIAL_DISTANCE_SUFFIX_RE = re.compile(
    r"\s*\(\s*\d+\s*m\s*\)\s*$",
    re.IGNORECASE,
)


def normalize_uploaded_athlete_name(value: object) -> str:
    """Remove the Meet Program AFL marker without otherwise changing a name."""
    return str(value).replace("(AFL)", "").strip()


def _clean_interprovincial_group_name(value: object) -> str:
    """Remove only the source format's trailing discipline-distance suffix."""
    return INTERPROVINCIAL_DISTANCE_SUFFIX_RE.sub("", str(value)).strip()


def _as_int(value):
    if value is None or value == "":
        return None
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


INTERPROVINCIAL_HEADERS = (
    "lane",
    "heat number",
    "group name",
    "athlete",
    "athlete number",
    "province",
)


def _is_interprovincial_layout(ws) -> bool:
    """Detect the supplied repeated-header format without relying on row numbers."""
    for row in ws.iter_rows(values_only=True):
        values = tuple(str(value or "").strip().casefold() for value in row[:6])
        if values == INTERPROVINCIAL_HEADERS:
            return True
    return False


def _parse_local_excel_records(ws) -> list[dict[str, Any]]:
    discipline = None
    current_heat = None
    records: list[dict[str, Any]] = []
    source_order = 0

    for row_num, row in enumerate(ws.iter_rows(values_only=True), start=1):
        a = row[0] if len(row) > 0 else None
        b = row[1] if len(row) > 1 else None
        c = row[2] if len(row) > 2 else None
        d = row[3] if len(row) > 3 else None
        s = str(a).strip() if a is not None else ""

        if s.lower() == "running heats":
            discipline = "running"
            current_heat = None
            continue
        if s.lower() == "swimming heats":
            discipline = "swimming"
            current_heat = None
            continue
        heat_match = HEAT_RE.match(s)
        if heat_match:
            current_heat = int(heat_match.group(1))
            continue
        if s == "#" or not s or discipline not in {"running", "swimming"}:
            continue

        athlete_number = _as_int(a)
        lane = _as_int(d)
        if athlete_number is None or b is None or c is None or current_heat is None:
            continue
        source_order += 1
        records.append(
            {
                "source_order": source_order,
                "discipline": discipline,
                "heat": current_heat,
                "athlete_number": str(athlete_number),
                "athlete_name": normalize_uploaded_athlete_name(b),
                "group_name": str(c).strip(),
                "province": None,
                "lane": lane,
                "row": row_num,
            }
        )
    return records


def _parse_interprovincial_excel_records(ws) -> list[dict[str, Any]]:
    discipline = None
    records: list[dict[str, Any]] = []
    source_order = 0

    for row_num, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = list(row) + [None] * max(0, 6 - len(row))
        lane_value, heat_value, group_value, name_value, number_value, province_value = values[:6]
        first_text = str(lane_value).strip() if lane_value is not None else ""
        if first_text.casefold() == "running heats":
            discipline = "running"
            continue
        if first_text.casefold() == "swimming heats":
            discipline = "swimming"
            continue
        if discipline not in {"running", "swimming"}:
            continue

        lane = _as_int(lane_value)
        heat = _as_int(heat_value)
        athlete_number = _as_int(number_value)
        if lane is None or heat is None or athlete_number is None or not group_value or not name_value:
            continue
        source_order += 1
        records.append(
            {
                "source_order": source_order,
                "discipline": discipline,
                "heat": heat,
                "athlete_number": str(athlete_number),
                "athlete_name": normalize_uploaded_athlete_name(name_value),
                "group_name": _clean_interprovincial_group_name(group_value),
                "province": str(province_value).strip() if province_value not in (None, "") else None,
                "lane": lane,
                "row": row_num,
            }
        )
    return records


def parse_master_entries(path_or_file) -> dict[str, Any]:
    """Parse an Excel or PDF Master Entries source into the canonical structure."""
    if isinstance(path_or_file, (str, Path)):
        suffix = Path(path_or_file).suffix.lower()
    else:
        suffix = Path(getattr(path_or_file, "name", "")).suffix.lower()
    if suffix == ".pdf":
        return parse_master_entries_pdf(path_or_file)

    wb = openpyxl.load_workbook(path_or_file, read_only=True, data_only=True)
    if not wb.sheetnames:
        raise ValueError("Workbook contains no worksheets.")
    ws = wb[wb.sheetnames[0]]

    try:
        is_interprovincial = _is_interprovincial_layout(ws)
        records = (
            _parse_interprovincial_excel_records(ws)
            if is_interprovincial
            else _parse_local_excel_records(ws)
        )
    finally:
        wb.close()

    if not records:
        raise ValueError("No athlete records were found in the Master Entries workbook.")

    # Build the canonical athlete list from the running roster ordering.
    run_records = [r for r in records if r["discipline"] == "running"]
    swim_records = [r for r in records if r["discipline"] == "swimming"]

    by_athlete: dict[str, dict[str, Any]] = {}
    for rec in run_records:
        aid = rec["athlete_number"]
        if aid in by_athlete:
            raise ValueError(f"Duplicate athlete number in running entries: {aid}")
        by_athlete[aid] = {
            "sort_order": len(by_athlete) + 1,
            "athlete_number": aid,
            "athlete_name": rec["athlete_name"],
            "group_name": rec["group_name"],
            "province": rec.get("province"),
            "running_heat": rec["heat"],
            "running_lane": rec["lane"],
            "swimming_heat": None,
            "swimming_lane": None,
        }

    missing_from_run = []
    for rec in swim_records:
        aid = rec["athlete_number"]
        if aid not in by_athlete:
            # Keep swimming-only entrants because TimeDrops needs them, but flag them later.
            by_athlete[aid] = {
                "sort_order": len(by_athlete) + 1,
                "athlete_number": aid,
                "athlete_name": rec["athlete_name"],
                "group_name": rec["group_name"],
                "province": rec.get("province"),
                "running_heat": None,
                "running_lane": None,
                "swimming_heat": rec["heat"],
                "swimming_lane": rec["lane"],
            }
            missing_from_run.append(aid)
            continue
        existing = by_athlete[aid]
        if existing["swimming_heat"] is not None:
            raise ValueError(f"Duplicate athlete number in swimming entries: {aid}")
        running_province = existing.get("province")
        swimming_province = rec.get("province")
        if running_province and swimming_province and running_province != swimming_province:
            raise ValueError(
                f"Athlete {aid} has conflicting Province abbreviations: "
                f"{running_province} in Running Heats and {swimming_province} in Swimming Heats."
            )
        if not running_province and swimming_province:
            existing["province"] = swimming_province
        existing["swimming_heat"] = rec["heat"]
        existing["swimming_lane"] = rec["lane"]

    return {
        "source_type": "xlsx_interprovincial" if is_interprovincial else "xlsx",
        "athletes": list(by_athlete.values()),
        "running_records": run_records,
        "swimming_records": swim_records,
        "running_heats": sorted({r["heat"] for r in run_records}),
        "swimming_heats": sorted({r["heat"] for r in swim_records}),
        "swimming_only_athletes": missing_from_run,
    }
