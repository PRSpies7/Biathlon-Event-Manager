"""Current category record references supplied by the user, separate from results."""
from decimal import Decimal, InvalidOperation
import re

import openpyxl

from parsers.season_results.common import athlete_number, clean_value, text

RECORD_HEADERS = ["Category", "Athlete number", "Athlete name", "Total points"]


def category_key(value):
    value = text(value).upper()
    value = re.sub(r"\bU\s*/?\s*0*(\d+)", r"U\1", value)
    value = re.sub(r"\b(?:JNR|JUNIOR)\b", "JUNIORS", value)
    value = re.sub(r"\bSENIOR\b", "SENIORS", value)
    value = re.sub(r"\bMASTER\b", "MASTERS", value)
    value = re.sub(r"^(BOYS|GIRLS)\s+(U\d+)\b", r"\2 \1", value)
    return value


def record_category_key(value):
    """Match category meaning, never fuzzy-match a neighbouring age or sex."""
    value = category_key(value)
    value = re.sub(r"\bUNDER\s*[-/]?\s*0*(\d+)\b", r"U\1", value)
    value = re.sub(r"\bU\s*[-/]?\s*0*(\d+)\b", r"U\1", value)
    tokens = re.findall(r"[A-Z]+|\d+", value)
    female = bool(set(tokens) & {"F", "FEMALE", "FEMALES", "GIRL", "GIRLS", "WOMAN", "WOMEN", "LADIES"})
    male = bool(set(tokens) & {"M", "MALE", "MALES", "BOY", "BOYS", "MAN", "MEN"})
    sex = "F" if female and not male else "M" if male and not female else ""
    if not sex:
        return value
    under = re.search(r"\bU(\d+)\b", value)
    if under:
        return f"U{int(under[1])} {sex}"
    if "MASTERS" in tokens:
        ages = re.findall(r"\d+", value)
        if len(ages) == 1:
            return f"MASTERS {int(ages[0])} {sex}"
    for group in ("JUNIORS", "SENIORS"):
        if group in tokens:
            return f"{group} {sex}"
    return value


def validate_record_benchmarks(records, *, allow_missing_numbers=True):
    clean, seen = [], set()
    for row in records:
        category = text(row.get("category"))
        number = athlete_number(row.get("athlete_number"))
        points = clean_value(row.get("total_points"))
        name = text(row.get("athlete_name"))
        if not number and not points and not name:
            continue  # Unfilled category rows in the template.
        if not category or (not number and not allow_missing_numbers) or not points:
            raise ValueError("Each record needs an age group and record points. Athlete details are optional.")
        key = record_category_key(category)
        if key in seen:
            raise ValueError(f"More than one current record was supplied for {category}.")
        try:
            value = Decimal(points)
            if not value.is_finite() or value < 0:
                raise InvalidOperation
        except InvalidOperation as exc:
            raise ValueError(f"Total points for {category} must be a non-negative number.") from exc
        seen.add(key)
        clean.append({"category_key": key, "category": category, "athlete_number": number,
                      "athlete_name": name, "total_points": points})
    if not clean:
        raise ValueError("No filled record references were found.")
    return clean


def parse_record_benchmarks(source):
    if hasattr(source, "seek"):
        source.seek(0)
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    try:
        layout = records_template_rows(workbook)
        if layout:
            return validate_record_benchmarks([
                {"category": sheet.cell(row, 2).value, "athlete_name": sheet.cell(row, 3).value,
                 "athlete_number": "", "total_points": str(Decimal(str(sheet.cell(row, 5).value)).quantize(Decimal("0.01")))}
                for sheet, row in layout], allow_missing_numbers=True)
        rows = iter(workbook.worksheets[0].values)
        header = next(rows, ())
        aliases = {"age group": "category", "age category": "category", "record points": "total points"}
        mapping = {aliases.get(text(v).casefold(), text(v).casefold()): i for i, v in enumerate(header)}
        if not {"category", "total points"}.issubset(mapping):
            raise ValueError("The records file needs Age group (or Category) and Record points (or Total points). Athlete details are optional.")
        records = [{label.replace(" ", "_"): row[i] if i < len(row) else None
                    for label, i in mapping.items() if label in {h.casefold() for h in RECORD_HEADERS}}
                   for row in rows]
    finally:
        workbook.close()
    return validate_record_benchmarks(records)


def records_template_rows(workbook):
    """Locate the supplied B:L records layout without relying on filename/year."""
    found = []
    for sheet in workbook:
        header = next((i for i in range(1, min(sheet.max_row, 20) + 1)
                       if text(sheet.cell(i, 2).value).upper() == "AGE GROUP"
                       and text(sheet.cell(i, 5).value).upper() == "RECORD POINTS"), None)
        if header:
            for i in range(header + 1, sheet.max_row + 1):
                if text(sheet.cell(i, 2).value) and sheet.cell(i, 5).value is not None:
                    found.append((sheet, i))
    return found
